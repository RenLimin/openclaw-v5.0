"""Word SDK — 基于 python-docx 的程序化构建与解析"""

from __future__ import annotations

import os
from datetime import datetime
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .exceptions import OfficeEngineError, OfficeParseError


class WordDocument:
    """Word 文档引擎

    统一接口：生命周期 + 元数据 + 解析 + 程序化构建
    """

    format = "docx"

    def __init__(self, path: Optional[str] = None):
        self._path = path
        self._doc = Document(path) if path and os.path.exists(path) else Document()
        self._closed = False
        # 默认字体
        self.set_font_default("Microsoft YaHei", 11)

    # ── 生命周期 ──────────────────────────────────────────────

    def save(self, path: Optional[str] = None) -> str:
        """保存文档，返回最终路径"""
        if self._closed:
            raise OfficeEngineError("document already closed")
        out = path or self._path
        if not out:
            raise OfficeEngineError("no save path specified")
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        self._doc.save(out)
        self._path = out
        return out

    def close(self):
        """关闭并释放资源"""
        self._closed = True
        self._doc = None

    # ── 元数据 ────────────────────────────────────────────────

    @property
    def metadata(self) -> dict:
        """作者/标题/创建时间等"""
        cp = self._doc.core_properties
        return {
            "author": cp.author,
            "title": cp.title,
            "subject": cp.subject,
            "keywords": cp.keywords,
            "created": cp.created.isoformat() if cp.created else None,
            "modified": cp.modified.isoformat() if cp.modified else None,
            "last_modified_by": cp.last_modified_by,
        }

    # ── 默认字体 ──────────────────────────────────────────────

    def set_font_default(self, font_name: str = "Microsoft YaHei", font_size: int = 11):
        """设置默认字体（含中文字体 eastAsia）"""
        style = self._doc.styles["Normal"]
        style.font.name = font_name
        style.font.size = Pt(font_size)
        # 中文字体必须同时设置 eastAsia
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            from docx.oxml import OxmlElement
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), font_name)
        rFonts.set(qn("w:ascii"), font_name)
        rFonts.set(qn("w:hAnsi"), font_name)

    # ── 构建 API ──────────────────────────────────────────────

    def add_heading(self, text: str, level: int = 1):
        """添加标题（1-9 级）"""
        if not (1 <= level <= 9):
            raise OfficeEngineError(f"heading level must be 1-9, got {level}")
        heading = self._doc.add_heading(text, level=level)
        # 标题也确保中文字体
        for run in heading.runs:
            self._set_run_font(run, "Microsoft YaHei")
        return heading

    def add_paragraph(
        self,
        text: str,
        style: Optional[str] = None,
        bold: bool = False,
        italic: bool = False,
        font_size: Optional[int] = None,
        color: Optional[str] = None,
        align: Optional[str] = None,
    ):
        """添加段落"""
        p = self._doc.add_paragraph(style=style) if style else self._doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if font_size:
            run.font.size = Pt(font_size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        if align:
            align_map = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }
            if align not in align_map:
                raise OfficeEngineError(f"unknown align: {align}")
            p.alignment = align_map[align]
        # 中文字体
        self._set_run_font(run, "Microsoft YaHei")
        return p

    def add_table(
        self,
        rows: int,
        cols: int,
        data: Optional[list[list]] = None,
        style: str = "Table Grid",
    ):
        """添加表格，data 为二维数组"""
        table = self._doc.add_table(rows=rows, cols=cols)
        try:
            table.style = style
        except KeyError:
            # 样式不存在时回退到 Table Grid
            table.style = "Table Grid"
        if data:
            for i, row_data in enumerate(data):
                if i >= rows:
                    break
                row = table.rows[i]
                for j, cell_text in enumerate(row_data):
                    if j >= cols:
                        break
                    cell = row.cells[j]
                    cell.text = str(cell_text)
                    # 确保中文字体
                    for p in cell.paragraphs:
                        for run in p.runs:
                            self._set_run_font(run, "Microsoft YaHei")
        return table

    def add_image(self, path: str, width: Optional[float] = None, height: Optional[float] = None):
        """添加图片，width/height 单位为英寸"""
        kwargs = {}
        if width:
            kwargs["width"] = Inches(width)
        if height:
            kwargs["height"] = Inches(height)
        return self._doc.add_picture(path, **kwargs)

    def add_list(self, items: list[str], ordered: bool = False):
        """添加有序/无序列表"""
        style = "List Number" if ordered else "List Bullet"
        for item in items:
            p = self._doc.add_paragraph(item, style=style)
            for run in p.runs:
                self._set_run_font(run, "Microsoft YaHei")
        return items

    def add_page_break(self):
        """添加分页符"""
        self._doc.add_page_break()

    def set_header(self, text: str):
        """设置页眉"""
        for section in self._doc.sections:
            header = section.header
            # 清空现有段落
            for p in list(header.paragraphs):
                p._element.getparent().remove(p._element)
            p = header.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                self._set_run_font(run, "Microsoft YaHei")

    def set_footer(self, text: str, page_number: bool = False):
        """设置页脚，可选页码"""
        for section in self._doc.sections:
            footer = section.footer
            for p in list(footer.paragraphs):
                p._element.getparent().remove(p._element)
            p = footer.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if text:
                run = p.add_run(text)
                self._set_run_font(run, "Microsoft YaHei")
            if page_number:
                if text:
                    p.add_run("  ")
                # 插入页码域
                from docx.oxml import OxmlElement
                run = p.add_run()
                fldChar1 = OxmlElement("w:fldChar")
                fldChar1.set(qn("w:fldCharType"), "begin")
                run._element.append(fldChar1)
                instrText = OxmlElement("w:instrText")
                instrText.set(qn("xml:space"), "preserve")
                instrText.text = "PAGE \\* MERGEFORMAT"
                run._element.append(instrText)
                fldChar2 = OxmlElement("w:fldChar")
                fldChar2.set(qn("w:fldCharType"), "end")
                run._element.append(fldChar2)
                self._set_run_font(run, "Microsoft YaHei")

    def find_replace(self, find_text: str, replace_text: str) -> int:
        """全局查找替换，返回替换次数"""
        count = 0
        # 遍历所有段落
        for p in self._doc.paragraphs:
            if find_text in p.text:
                # 简单替换：重建 run
                full_text = p.text
                new_text = full_text.replace(find_text, replace_text)
                count += full_text.count(find_text)
                # 清空现有 runs
                for run in list(p.runs):
                    run._element.getparent().remove(run._element)
                run = p.add_run(new_text)
                self._set_run_font(run, "Microsoft YaHei")
        # 遍历所有表格单元格
        for table in self._doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if find_text in p.text:
                            full_text = p.text
                            new_text = full_text.replace(find_text, replace_text)
                            count += full_text.count(find_text)
                            for run in list(p.runs):
                                run._element.getparent().remove(run._element)
                            run = p.add_run(new_text)
                            self._set_run_font(run, "Microsoft YaHei")
        return count

    # ── 解析 API ──────────────────────────────────────────────

    def parse(self) -> dict:
        """解析文档，返回结构化 JSON"""
        result = {
            "format": "docx",
            "metadata": self.metadata,
            "headings": [],
            "paragraphs": [],
            "tables": [],
            "images": [],
            "sections": [],
        }

        for i, p in enumerate(self._doc.paragraphs):
            style_name = p.style.name if p.style else "Normal"
            text = p.text
            if not text and style_name != "Normal":
                continue

            # 识别标题
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.replace("Heading", "").strip())
                except ValueError:
                    level = 1
                result["headings"].append({
                    "index": i,
                    "level": level,
                    "text": text,
                })
            elif style_name in ("List Bullet", "List Number"):
                result["paragraphs"].append({
                    "index": i,
                    "type": "list",
                    "ordered": style_name == "List Number",
                    "text": text,
                    "style": style_name,
                })
            else:
                result["paragraphs"].append({
                    "index": i,
                    "type": "paragraph",
                    "text": text,
                    "style": style_name,
                })

        # 表格
        for t_idx, table in enumerate(self._doc.tables):
            rows_data = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                rows_data.append(row_data)
            result["tables"].append({
                "index": t_idx,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "data": rows_data,
            })

        # 图片（粗略统计 inline shapes）
        inline_shapes = self._doc.inline_shapes
        result["images"] = [
            {"index": i, "type": shape.type.__name__ if hasattr(shape.type, "__name__") else str(shape.type)}
            for i, shape in enumerate(inline_shapes)
        ]

        # section 信息
        for s_idx, section in enumerate(self._doc.sections):
            result["sections"].append({
                "index": s_idx,
                "page_width": section.page_width.inches if section.page_width else None,
                "page_height": section.page_height.inches if section.page_height else None,
                "header_text": section.header.paragraphs[0].text if section.header.paragraphs else "",
                "footer_text": section.footer.paragraphs[0].text if section.footer.paragraphs else "",
            })

        return result

    def to_markdown(self) -> str:
        """转换为 Markdown"""
        lines = []

        for p in self._doc.paragraphs:
            style_name = p.style.name if p.style else "Normal"
            text = p.text

            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.replace("Heading", "").strip())
                except ValueError:
                    level = 1
                lines.append(f"{'#' * level} {text}")
                lines.append("")
            elif style_name == "List Bullet":
                lines.append(f"- {text}")
            elif style_name == "List Number":
                # 简单处理，都用数字 1.
                lines.append(f"1. {text}")
            elif text.strip():
                lines.append(text)
                lines.append("")
            else:
                lines.append("")

        # 表格
        for table in self._doc.tables:
            if not table.rows:
                continue
            # 表头
            header = [cell.text for cell in table.rows[0].cells]
            lines.append("| " + " | ".join(header) + " |")
            lines.append("| " + " | ".join(["---"] * len(header)) + " |")
            # 数据行
            for row in table.rows[1:]:
                row_data = [cell.text for cell in row.cells]
                lines.append("| " + " | ".join(row_data) + " |")
            lines.append("")

        return "\n".join(lines).rstrip() + "\n"

    # ── 内部工具 ──────────────────────────────────────────────

    @staticmethod
    def _set_run_font(run, font_name: str):
        """设置 run 的字体（含 eastAsia）"""
        run.font.name = font_name
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            from docx.oxml import OxmlElement
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:eastAsia"), font_name)
        rFonts.set(qn("w:ascii"), font_name)
        rFonts.set(qn("w:hAnsi"), font_name)
