"""pytest 配置：注入组件根与 L3 组件根到 sys.path。

- 组件根：使 ``contract_review`` / ``repository`` 可导入。
- L3 contract-management 组件根：使 ``contract_management`` 可导入
  （依赖方向 L4→L3，参考 L3 tests/conftest 模式，不修改 L3）。
另提供真实文件夹具构造器（PDF/docx/扫描件），不使用 mock。
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
COMPONENT_ROOT = os.path.dirname(os.path.dirname(HERE))
REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(COMPONENT_ROOT)))
L3_ROOT = os.path.join(
    REPO_ROOT, "L3-business", "components", "contract-management"
)

for path in (COMPONENT_ROOT, L3_ROOT):
    if path not in sys.path:
        sys.path.insert(0, path)


# ------------------------------------------------------------
# 真实文件夹具构造器
# ------------------------------------------------------------

def make_native_pdf(path: str, text: str = "CONTRACT NATIVE TEXT 2026") -> str:
    """生成含真实文本层的 PDF（pymupdf 写入，pypdf 可抽取）。"""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontname="helv", fontsize=16)
    doc.save(path)
    doc.close()
    return path


def make_docx(path: str) -> str:
    """生成含段落与表格的真实 docx。"""
    import docx

    document = docx.Document()
    document.add_paragraph("技术服务合同 DOCX NATIVE")
    document.add_paragraph("甲方：测试公司")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "序号"
    table.cell(0, 1).text = "标的"
    table.cell(0, 2).text = "金额"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "安全测评服务"
    table.cell(1, 2).text = "10000"
    document.save(path)
    return path


def make_text_image(path: str, text: str = "CONTRACT 2026") -> str:
    """生成含文字的真实图片（供 OCR）。"""
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (900, 220), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(
        "/System/Library/Fonts/Supplemental/Arial.ttf", 72
    )
    draw.text((40, 70), text, fill="black", font=font)
    img.save(path)
    return path


def make_scanned_pdf(path: str, text: str = "CONTRACT 2026") -> str:
    """生成无文本层的扫描型 PDF（页面是一张文字图片）。"""
    import fitz
    from PIL import Image, ImageDraw, ImageFont
    import io

    img = Image.new("RGB", (900, 220), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(
        "/System/Library/Fonts/Supplemental/Arial.ttf", 72
    )
    draw.text((40, 70), text, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(0, 0, 612, 150), stream=buf.getvalue())
    doc.save(path)
    doc.close()
    return path
