"""报告测试 — VERIFICATION：双格式结构 + 免责声明。"""
import os

import pytest

from contract_review.kb_match import KbReview, KbMatchResult
from contract_review.report import (
    DISCLAIMER, build_markdown, write_markdown, build_excel,
)


@pytest.fixture
def risk_dict():
    return {
        "overall_risk": "medium",
        "summary": {"pass": 18, "warning": 3, "fail": 0},
        "scan_time": "2026-10-28T00:00:00",
        "findings": [
            {"id": "1.5", "category": "主体信息", "item": "签字盖章",
             "status": "pass", "risk": "high", "law": "§490"},
            {"id": "9.2", "category": "保密条款", "item": "保密期限",
             "status": "warning", "risk": "medium", "law": "—"},
        ],
    }


@pytest.fixture
def kb_review():
    return KbReview(
        satisfaction="partial",
        referenced=True,
        items=[KbMatchResult(input_name="安全测评服务", matched=True, score=0.9,
                             product="安全测评服务", sku="SKU-1",
                             alert="低匹配度，需人工确认")],
    )


def test_markdown_structure(risk_dict, kb_review):
    md = build_markdown(contract_ref="C1", source="upload",
                        risk=risk_dict, kb=kb_review)
    assert "# 合同智能审查报告" in md
    assert "综合风险" in md and "medium" in md
    assert "产品满足度" in md
    assert "1.5" in md and "§490" in md
    assert DISCLAIMER in md
    assert "部分满足" in md


def test_markdown_unreferenced_kb(risk_dict):
    kb = KbReview(satisfaction="unsatisfied", referenced=False)
    md = build_markdown(contract_ref="C1", source="upload", risk=risk_dict, kb=kb)
    assert "未做产品参照" in md


def test_markdown_no_items(risk_dict):
    kb = KbReview(satisfaction="unsatisfied", referenced=True, items=[])
    md = build_markdown(contract_ref="C1", source="upload", risk=risk_dict, kb=kb)
    assert "未从合同提取到可参照标的" in md


def test_write_markdown(tmp_path, risk_dict, kb_review):
    out = tmp_path / "report.md"
    path = write_markdown(contract_ref="C1", source="upload",
                          risk=risk_dict, kb=kb_review, out_path=str(out))
    assert os.path.isfile(path)
    content = open(path, encoding="utf-8").read()
    assert DISCLAIMER in content


def test_excel_generation(tmp_path, risk_dict, kb_review):
    import openpyxl
    out = tmp_path / "report.xlsx"
    path = build_excel(contract_ref="C1", source="upload",
                       risk=risk_dict, kb=kb_review, out_path=str(out))
    assert os.path.isfile(path)
    wb = openpyxl.load_workbook(path)
    assert set(wb.sheetnames) == {"概览", "风险明细", "KB匹配"}
    ws = wb["风险明细"]
    assert ws.cell(row=2, column=1).value == "1.5"
    ws3 = wb["KB匹配"]
    assert ws3.cell(row=2, column=1).value == "安全测评服务"


def test_excel_requires_openpyxl(tmp_path, risk_dict, kb_review, monkeypatch):
    import builtins
    real_import = builtins.__import__

    def fake_import(name, *a, **k):
        if name == "openpyxl":
            raise ImportError("no openpyxl")
        return real_import(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(ImportError):
        build_excel(contract_ref="C1", source="upload",
                    risk=risk_dict, kb=kb_review,
                    out_path=str(tmp_path / "r.xlsx"))
