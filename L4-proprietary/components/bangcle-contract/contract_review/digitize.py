"""数字化（digitize）— PDF/docx 直读，扫描件走 L2 OCR。

按 DESIGN-DETAIL-v1.0 §3 的 DigitizedText 数据契约，兜底口径见 §6：
- PDF：优先文本层直读；无文本层（扫描件）→ L2 OCR。
- docx：文本 / 表格直读。
- 图片：走 L2 OCR。

依赖策略：
- pypdf / python-docx 为可选依赖，缺失时优雅报错（不引入硬依赖）。
- L2 OCR 通过 :class:`OCRAdapter` 适配；接口不确定时以清晰占位实现，
  并在 ``method`` 上如实标注。合同原文全程本地，不外发。
"""

from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Optional


class DigitizeError(RuntimeError):
    """数字化失败：缺库 / OCR 不可用 / 文件不可读等。"""


@dataclass
class DigitizedText:
    """数字化结果（DETAIL §3）。"""

    full_text: str
    method: str                          # "native_pdf" | "docx" | "ocr"
    ocr_confidence: Optional[float] = None


# ------------------------------------------------------------
# 直读实现
# ------------------------------------------------------------

def read_pdf_text(path: str) -> str:
    """用 pypdf 直读 PDF 文本层；无文本层返回空串。缺库时 DigitizeError。"""
    try:
        from pypdf import PdfReader
    except Exception as exc:  # pragma: no cover - 环境缺库路径
        raise DigitizeError(
            "PDF 直读需要可选依赖 pypdf（pip install pypdf）"
        ) from exc

    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            # 单页解析失败不致命，按空文本处理
            parts.append("")
    return "\n".join(parts).strip()


def read_docx_text(path: str) -> str:
    """用 python-docx 直读 docx 段落与表格。缺库时 DigitizeError。"""
    try:
        import docx  # python-docx
    except Exception as exc:  # pragma: no cover - 环境缺库路径
        raise DigitizeError(
            "docx 直读需要可选依赖 python-docx（pip install python-docx）"
        ) from exc

    document = docx.Document(path)
    lines = [para.text for para in document.paragraphs]

    # 表格：逐行逐单元格读取，保持表格文本不丢
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " ".join(c for c in cells if c)
            if line:
                lines.append(line)

    return "\n".join(line for line in lines if line is not None).strip()


# ------------------------------------------------------------
# OCR 适配（L2 OCR-001）
# ------------------------------------------------------------

class OCRAdapter:
    """L2 OCR-001 适配器。

    L2 OCR 的对外接口在不同文档中存在差异（skill 文档称
    ``digitalize_document_v5``，实际组件入口为
    ``digitalize_document``），故这里做**防御式适配**：

    - 通过 sys.path 注入 L2 OCR 组件 scripts 目录（依赖方向 L4→L2，
      不修改 L2）；
    - 优先解析实际存在的入口函数；两者都拿不到时抛 DigitizeError，
      调用方据此把合同标记为"未数字化"（DETAIL §9）。

    本适配器只在本地调用 L2，不发起任何公网网络请求。
    """

    def __init__(self, l2_scripts_dir: Optional[str] = None) -> None:
        self.l2_scripts_dir = l2_scripts_dir or self._default_l2_scripts_dir()

    @staticmethod
    def _default_l2_scripts_dir() -> str:
        # 组件根/../../.. -> 仓库根；再进 L2 OCR scripts
        component_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        repo_root = os.path.dirname(os.path.dirname(os.path.dirname(component_root)))
        return os.path.join(
            repo_root,
            "L2-infra", "components", "ocr-digitalization", "scripts",
        )

    def _load_entry(self):
        scripts_dir = self.l2_scripts_dir
        if not os.path.isdir(scripts_dir):
            raise DigitizeError(f"L2 OCR scripts 目录不存在: {scripts_dir}")
        import sys

        if scripts_dir not in sys.path:
            sys.path.insert(0, scripts_dir)
        try:
            import ocr_engine  # L2 包
        except Exception as exc:
            raise DigitizeError(f"无法导入 L2 ocr_engine: {exc}") from exc

        # 实际入口 digitalize_document（engine.py）
        entry = getattr(ocr_engine, "digitalize_document", None)
        if entry is None:
            # 兼容旧 skill 文档所述入口
            from ocr_engine import document as _doc  # type: ignore

            entry = getattr(_doc, "digitalize_document_v5", None)
        if entry is None:
            raise DigitizeError(
                "L2 OCR 入口不可用（digitalize_document / digitalize_document_v5 均缺失）"
            )
        return entry

    def recognize(self, path: str) -> DigitizedText:
        """对扫描件 / 图片执行 OCR，返回 method='ocr' 的 DigitizedText。"""
        entry = self._load_entry()
        try:
            # 显式指定 rapidocr：L2 auto 发现存在 lang 码不匹配
            # （auto 用 chi_sim+eng，rapidocr 后端声明 ch+en），
            # 显式后端可绕过该 L2 问题且不改 L2。
            try:
                result = entry(path, backend="rapidocr")
            except TypeError:
                # 兼容签名不同的入口（如旧 digitalize_document_v5）
                result = entry(path)
        except Exception as exc:
            raise DigitizeError(f"L2 OCR 识别失败: {exc}") from exc

        text = self._extract_text(result)
        confidence = self._extract_confidence(result)
        return DigitizedText(
            full_text=text,
            method="ocr",
            ocr_confidence=confidence,
        )

    @staticmethod
    def _extract_text(result) -> str:
        for attr in ("to_text",):
            fn = getattr(result, attr, None)
            if callable(fn):
                try:
                    return (fn() or "").strip()
                except Exception:
                    pass
        text = getattr(result, "text", None)
        if isinstance(text, str):
            return text.strip()
        return str(result).strip() if result is not None else ""

    @staticmethod
    def _extract_confidence(result) -> Optional[float]:
        # OCRResult.confidence 在 L2 类型里可能是百分制（0-100）
        conf = getattr(result, "confidence", None)
        try:
            value = float(conf)
        except (TypeError, ValueError):
            return None
        return value / 100.0 if value > 1.0 else value


# ------------------------------------------------------------
# 统一数字化入口
# ------------------------------------------------------------

def digitize(
    path: str,
    file_type: Optional[str] = None,
    ocr_adapter: Optional[OCRAdapter] = None,
) -> DigitizedText:
    """把本地合同文件数字化为全文文本。

    Args:
        path: 本地文件路径。
        file_type: 显式类型（pdf/docx/image）；缺省按扩展名判定。
        ocr_adapter: 可注入的 OCR 适配器（测试用）。

    Returns:
        DigitizedText
    """
    if not os.path.isfile(path):
        raise DigitizeError(f"文件不存在: {path}")

    if file_type is None:
        from .ingest import detect_file_type

        file_type = detect_file_type(path)

    if file_type == "pdf":
        text = read_pdf_text(path)
        if text:
            return DigitizedText(full_text=text, method="native_pdf")
        # 无文本层 → 扫描件，走 OCR
        adapter = ocr_adapter or OCRAdapter()
        return adapter.recognize(path)

    if file_type == "docx":
        return DigitizedText(full_text=read_docx_text(path), method="docx")

    if file_type == "image":
        adapter = ocr_adapter or OCRAdapter()
        return adapter.recognize(path)

    raise DigitizeError(f"不支持数字化的文件类型: {file_type}")
