"""
BDMS v2 — 额外引擎测试（合同引擎 + 工时引擎 + Excel写入层）

目标：提升覆盖率到 80%+
覆盖：
- contract_engine：合同编号校准、OA合同加载、合同信息关联
- hours_engine：工时加载、按项目/合同汇总、人天换算、POC工时写入
- build_v2_report：ExcelWriter 基本功能、导出函数、POC专用计算列
"""

import pytest
import pandas as pd
import numpy as np
from pathlib import Path
import sys
import tempfile

V2_DIR = Path(__file__).parent.parent / "src" / "delivery_center" / "v2"
sys.path.insert(0, str(V2_DIR))
sys.path.insert(0, str(V2_DIR / "engines"))
sys.path.insert(0, str(V2_DIR / "generators"))
sys.path.insert(0, str(V2_DIR / "config"))


# ============================================================
# 合同引擎测试
# ============================================================

class TestContractEngine:
    """合同引擎测试"""

    def test_calibrate_contract_no_normal(self):
        """合同编号校准：正常编号返回大写"""
        from contract_engine import calibrate_contract_no

        assert calibrate_contract_no("ht-001") == "HT-001"
        assert calibrate_contract_no("BC-2026-001") == "BC-2026-001"

    def test_calibrate_contract_no_strip_ampersand(self):
        """合同编号校准：去除 & 后面的内容"""
        from contract_engine import calibrate_contract_no

        assert calibrate_contract_no("HT-001&补充协议") == "HT-001"
        assert calibrate_contract_no("bc-002&v2") == "BC-002"

    def test_calibrate_contract_no_strip_spaces(self):
        """合同编号校准：去前后空格"""
        from contract_engine import calibrate_contract_no

        assert calibrate_contract_no("  HT-001  ") == "HT-001"

    def test_calibrate_contract_no_empty(self):
        """合同编号校准：空值返回空字符串"""
        from contract_engine import calibrate_contract_no

        assert calibrate_contract_no("") == ""
        assert calibrate_contract_no(None) == ""
        assert calibrate_contract_no(123) == ""  # 非字符串

    def test_load_oa_contracts_none_path(self):
        """OA合同加载：None 路径返回空结构 DataFrame"""
        from contract_engine import load_oa_contracts

        df = load_oa_contracts(None)
        assert len(df) == 0
        assert "合同编号" in df.columns
        assert "签约金额" in df.columns
        assert "合同归档日期" in df.columns
        assert "客户名称" in df.columns

    def test_load_oa_contracts_nonexistent_path(self):
        """OA合同加载：不存在的路径返回空结构"""
        from contract_engine import load_oa_contracts

        df = load_oa_contracts(Path("/nonexistent/path.xlsx"))
        assert len(df) == 0
        assert "合同编号" in df.columns

    def test_join_contract_info_empty_contracts(self):
        """合同关联：空合同表时补空列"""
        from contract_engine import join_contract_info

        df_projects = pd.DataFrame({
            "销售合同编号": ["HT-001", "HT-002"],
            "项目名称": ["项目A", "项目B"],
        })
        df_contracts = pd.DataFrame(columns=["合同编号", "签约金额"])

        result = join_contract_info(df_projects, df_contracts)

        # 保留所有原行
        assert len(result) == 2
        # 补了 oa_ 前缀列
        assert "oa_签约金额" in result.columns
        assert "oa_合同归档日期" in result.columns
        assert (result["oa_签约金额"] == "").all()

    def test_join_contract_info_basic_join(self):
        """合同关联：正常左连接，匹配正确"""
        from contract_engine import join_contract_info

        df_projects = pd.DataFrame({
            "销售合同编号": ["HT-001", "HT-002", "HT-003"],
            "项目名称": ["项目A", "项目B", "项目C"],
        })
        df_contracts = pd.DataFrame({
            "合同编号": ["HT-001", "HT-002"],
            "签约金额": [100000, 200000],
            "客户名称": ["客户A", "客户B"],
            "合同归档日期": ["2026-01-15", "2026-02-20"],
        })

        result = join_contract_info(df_projects, df_contracts)

        assert len(result) == 3
        # HT-001 匹配
        assert result.iloc[0]["oa_签约金额"] == 100000
        assert result.iloc[0]["oa_客户名称"] == "客户A"
        # HT-003 未匹配，应该是 NaN
        assert pd.isna(result.iloc[2]["oa_签约金额"])

    def test_join_contract_info_case_insensitive(self):
        """合同关联：项目侧大小写不敏感（项目侧会校准，合同侧假设已校准为大写）"""
        from contract_engine import join_contract_info

        df_projects = pd.DataFrame({
            "销售合同编号": ["ht-001", "ht-002"],  # 小写
            "项目名称": ["项目A", "项目B"],
        })
        df_contracts = pd.DataFrame({
            "合同编号": ["HT-001", "HT-002"],  # OA侧是大写（load_oa_contracts 校准后）
            "签约金额": [100, 200],
            "客户名称": ["客户A", "客户B"],
        })

        result = join_contract_info(df_projects, df_contracts)

        # 项目侧小写会被校准为大写，匹配成功
        assert result.iloc[0]["oa_签约金额"] == 100
        assert result.iloc[1]["oa_签约金额"] == 200

    def test_join_contract_info_ampersand(self):
        """合同关联：带 & 的合同号也能匹配"""
        from contract_engine import join_contract_info

        df_projects = pd.DataFrame({
            "销售合同编号": ["HT-001&补充", "HT-002"],
            "项目名称": ["项目A", "项目B"],
        })
        df_contracts = pd.DataFrame({
            "合同编号": ["HT-001", "HT-002"],
            "签约金额": [100, 200],
            "客户名称": ["客户A", "客户B"],
        })

        result = join_contract_info(df_projects, df_contracts)

        # HT-001&补充 校准后是 HT-001，应匹配
        assert result.iloc[0]["oa_签约金额"] == 100

    def test_add_contract_columns_wrapper(self):
        """add_contract_columns 封装函数可用"""
        try:
            from contract_engine import add_contract_columns

            df = pd.DataFrame({"销售合同编号": ["HT-001"]})
            contracts = pd.DataFrame(columns=["合同编号"])
            result = add_contract_columns(df, df_contracts=contracts)
            assert len(result) == 1
        except ImportError:
            pytest.skip("add_contract_columns not available in v2")


# ============================================================
# 工时引擎测试
# ============================================================

class TestHoursEngine:
    """工时引擎测试"""

    def test_load_workhour_detail_none_path(self):
        """工时加载：None 路径返回空结构"""
        from hours_engine import load_workhour_detail

        df = load_workhour_detail(None)
        assert len(df) == 0
        assert "项目名称" in df.columns
        assert "登记工时" in df.columns

    def test_load_workhour_pivot_none_path(self):
        """工时pivot加载：None 路径返回空结构"""
        from hours_engine import load_workhour_pivot

        df = load_workhour_pivot(None)
        assert len(df) == 0
        assert "项目编号" in df.columns
        assert "工时合计" in df.columns

    def test_aggregate_by_project_basic(self):
        """按项目汇总工时"""
        from hours_engine import aggregate_by_project

        df_detail = pd.DataFrame({
            "项目名称": ["项目A", "项目A", "项目B", "项目C"],
            "登记工时": ["8", "4", "6", "10"],
        })

        result = aggregate_by_project(df_detail)

        assert len(result) == 3
        # 项目A 合计 12 小时
        a_row = result[result["项目编号"] == "项目A"]
        assert a_row.iloc[0]["工时合计"] == 12.0
        # 按工时降序排列
        assert result.iloc[0]["工时合计"] >= result.iloc[-1]["工时合计"]

    def test_aggregate_by_project_empty(self):
        """按项目汇总：空数据返回空结果"""
        from hours_engine import aggregate_by_project

        result = aggregate_by_project(pd.DataFrame())
        assert len(result) == 0

    def test_aggregate_by_project_missing_columns(self):
        """按项目汇总：缺少必要列返回空"""
        from hours_engine import aggregate_by_project

        df = pd.DataFrame({"随机列": ["a", "b"]})
        result = aggregate_by_project(df)
        assert len(result) == 0

    def test_aggregate_by_contract_basic(self):
        """按合同汇总工时"""
        from hours_engine import aggregate_by_contract

        df_detail = pd.DataFrame({
            "合同编号": ["HT-001", "HT-001", "HT-002"],
            "登记工时": ["8", "16", "10"],
        })

        result = aggregate_by_contract(df_detail)

        assert len(result) == 2
        h1_row = result[result["合同编号"] == "HT-001"]
        assert h1_row.iloc[0]["工时合计"] == 24.0

    def test_aggregate_by_contract_empty(self):
        """按合同汇总：空数据返回空"""
        from hours_engine import aggregate_by_contract

        result = aggregate_by_contract(pd.DataFrame())
        assert len(result) == 0

    def test_hours_to_person_days_normal(self):
        """人天换算：正常换算"""
        from hours_engine import hours_to_person_days

        assert hours_to_person_days(8) == 1.0
        assert hours_to_person_days(16) == 2.0
        assert hours_to_person_days(4) == 0.5

    def test_hours_to_person_days_zero(self):
        """人天换算：0 和 None/NaN 返回 0.0"""
        from hours_engine import hours_to_person_days

        assert hours_to_person_days(0) == 0.0
        assert hours_to_person_days(None) == 0.0
        assert hours_to_person_days(np.nan) == 0.0

    def test_hours_to_person_days_custom_hours(self):
        """人天换算：自定义工时/天"""
        from hours_engine import hours_to_person_days

        assert hours_to_person_days(7, hours_per_day=7) == 1.0
        assert hours_to_person_days(14, hours_per_day=7) == 2.0

    def test_hours_to_person_days_round_2_decimals(self):
        """人天换算：保留 2 位小数"""
        from hours_engine import hours_to_person_days

        result = hours_to_person_days(10)
        assert round(result * 100) == result * 100  # 2位小数

    def test_add_poc_hours_column_normal(self):
        """POC工时写入：正常匹配写入"""
        from hours_engine import add_poc_hours_column

        df_poc = pd.DataFrame({
            "销售合同编号": ["HT-001", "HT-002", "HT-003"],
            "项目名称": ["项目A", "项目B", "项目C"],
        })
        df_hours = pd.DataFrame({
            "合同编号": ["HT-001", "HT-002"],
            "工时合计": [100.0, 200.0],
        })

        result = add_poc_hours_column(df_poc, df_hours)

        assert "POC项目工时合计（小时）" in result.columns
        assert result.iloc[0]["POC项目工时合计（小时）"] == 100.0
        assert result.iloc[1]["POC项目工时合计（小时）"] == 200.0
        assert result.iloc[2]["POC项目工时合计（小时）"] == 0.0  # 未匹配

    def test_add_poc_hours_column_empty_hours(self):
        """POC工时写入：空工时表填 0"""
        from hours_engine import add_poc_hours_column

        df_poc = pd.DataFrame({"销售合同编号": ["HT-001"]})
        df_hours = pd.DataFrame(columns=["合同编号", "工时合计"])

        result = add_poc_hours_column(df_poc, df_hours)
        assert result.iloc[0]["POC项目工时合计（小时）"] == 0.0

    def test_add_poc_hours_column_case_insensitive(self):
        """POC工时写入：POC侧小写合同号也能匹配工时表的大写编号"""
        from hours_engine import add_poc_hours_column

        df_poc = pd.DataFrame({"销售合同编号": ["ht-001", "ht-002"]})  # 小写
        df_hours = pd.DataFrame({
            "合同编号": ["HT-001", "HT-002"],  # 大写
            "工时合计": [50.0, 75.0],
        })

        result = add_poc_hours_column(df_poc, df_hours)
        # POC侧校准后大写，匹配成功
        assert result.iloc[0]["POC项目工时合计（小时）"] == 50.0
        assert result.iloc[1]["POC项目工时合计（小时）"] == 75.0

    def test_add_person_days_column(self):
        """人天列添加：正确换算"""
        from hours_engine import add_person_days_column

        df = pd.DataFrame({"POC项目工时合计（小时）": [8, 16, 0]})
        result = add_person_days_column(df)

        assert "POC项目工时合计（人天）" in result.columns
        assert result.iloc[0]["POC项目工时合计（人天）"] == 1.0
        assert result.iloc[1]["POC项目工时合计（人天）"] == 2.0
        assert result.iloc[2]["POC项目工时合计（人天）"] == 0.0

    def test_add_person_days_column_missing_hours(self):
        """人天列添加：缺少工时列时填 0"""
        from hours_engine import add_person_days_column

        df = pd.DataFrame({"其他列": [1, 2]})
        result = add_person_days_column(df)
        assert result.iloc[0]["POC项目工时合计（人天）"] == 0.0

    def test_get_poc_project_hours_empty(self):
        """便捷入口：空文件返回空结构"""
        from hours_engine import get_poc_project_hours_by_contract

        result = get_poc_project_hours_by_contract(None)
        assert len(result) == 0
        assert "合同编号" in result.columns

    def test_load_workhour_pivot_real_excel(self):
        """工时pivot加载：真实 Excel 文件 + 列名规范化 + 总计行过滤"""
        from hours_engine import load_workhour_pivot

        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "workhour_pivot.xlsx"
            df = pd.DataFrame({
                "行标签": ["项目A", "项目B", "项目C", "总计"],
                "求和项:登记工时": ["100.5", "200", "150.5", "451"],
            })
            with pd.ExcelWriter(xlsx_path) as writer:
                df.to_excel(writer, index=False, sheet_name="按项目汇总")

            result = load_workhour_pivot(xlsx_path, header_row=0)

            # 3个项目 + 总计行被过滤
            assert len(result) == 3
            # 列名被规范化
            assert list(result.columns) == ["项目编号", "工时合计"]
            # 工时是数值
            assert result.iloc[0]["工时合计"] == 100.5
            assert result["工时合计"].dtype in [np.float64, np.int64, float]

    def test_load_workhour_pivot_nonexistent(self):
        """工时pivot加载：文件不存在返回空结构"""
        from hours_engine import load_workhour_pivot

        result = load_workhour_pivot(Path("/nonexistent.xlsx"))
        assert len(result) == 0
        assert "项目编号" in result.columns

    def test_load_workhour_pivot_bad_columns(self):
        """工时pivot加载：列名不匹配返回空结构"""
        from hours_engine import load_workhour_pivot

        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "bad_pivot.xlsx"
            df = pd.DataFrame({"奇怪的列名": [1, 2], "另一列": [3, 4]})
            with pd.ExcelWriter(xlsx_path) as writer:
                df.to_excel(writer, index=False, sheet_name="按项目汇总")

            result = load_workhour_pivot(xlsx_path, header_row=0)
            assert len(result) == 0

    def test_load_workhour_detail_real_excel(self):
        """工时明细加载：真实 Excel 文件"""
        from hours_engine import load_workhour_detail

        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "workhour_detail.xlsx"
            df = pd.DataFrame({
                "项目名称": ["项目A", "项目A", "项目B"],
                "合同编号": ["HT-001", "HT-001", "HT-002"],
                "登记工时": ["8", "4", "6"],
                "人员姓名": ["张三", "李四", "王五"],
            })
            df.to_excel(xlsx_path, index=False, sheet_name="工时填报情况查询")

            result = load_workhour_detail(xlsx_path, header_row=0)
            assert len(result) == 3
            assert "项目名称" in result.columns
            assert "登记工时" in result.columns


# ============================================================
# POC 专用计算列测试（build_v2_report 中的 _add_poc_specific_columns）
# ============================================================

class TestPocSpecificColumns:
    """POC 专用计算列测试"""

    def test_poc_duration_columns_added(self):
        """POC 专用列：持续周期相关列被添加"""
        from build_v2_report import ReportConfig, ReportEngineOrchestrator

        cfg = ReportConfig(report_month="202606")
        orch = ReportEngineOrchestrator(cfg)

        df = pd.DataFrame({
            "所属项目": ["POC-A", "POC-B", "POC-C"],
            "立项日期": ["2026-01-10", "2026-03-15", "2026-05-20"],
            "实际结项日期": ["2026-03-01", "2026-05-01", ""],
            "销售合同编号": ["HT-001", "HT-002", ""],
            "合同归档日期": ["2026-02-01", "", ""],
            "项目类型(概览)": ["POC", "提前实施", "POC"],
        })

        result = orch._add_poc_specific_columns(df.copy())

        # 至少新增了持续周期相关列
        assert len(result.columns) > len(df.columns)
        duration_cols = [c for c in result.columns if "持续周期" in c]
        assert len(duration_cols) >= 2  # 至少有天数和统计两列

    def test_poc_duration_with_actual_end_date(self):
        """POC 持续周期：有实际结项日期时计算正确"""
        from build_v2_report import ReportConfig, ReportEngineOrchestrator

        cfg = ReportConfig(report_month="202606")
        orch = ReportEngineOrchestrator(cfg)

        df = pd.DataFrame({
            "所属项目": ["POC-1", "POC-2"],
            "立项日期": ["2026-01-01", "2026-03-01"],
            "实际结项日期": ["2026-01-31", ""],  # 30天 / 空
            "销售合同编号": ["HT-001", "HT-002"],
            "合同归档日期": ["2026-02-01", ""],
            "项目类型(概览)": ["提前实施", "POC"],
        })

        result = orch._add_poc_specific_columns(df.copy())

        # 找到持续周期天数列
        day_cols = [c for c in result.columns if "持续周期" in c and "天" in c]
        assert len(day_cols) > 0, "应包含持续周期天数列"
        
        # 第一个项目有实际结项日期，应有值
        col_name = day_cols[0]
        assert pd.notna(result.iloc[0][col_name])
        # 第二个项目空日期，可能为 NaN 或 0


# ============================================================
# ExcelWriter 测试
# ============================================================

class TestExcelWriter:
    """Excel 写入器测试"""

    def test_writer_creates_file(self):
        """ExcelWriter 创建输出文件"""
        from build_v2_report import ExcelWriter

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "test_report.xlsx"
            writer = ExcelWriter(output_path)

            df = pd.DataFrame({"A": [1, 2, 3], "B": ["x", "y", "z"]})
            writer.write_dataframe(df, "测试Sheet")
            writer.save()

            assert output_path.exists()
            assert output_path.stat().st_size > 0

    def test_writer_multiple_sheets(self):
        """ExcelWriter 写入多个 Sheet"""
        from openpyxl import load_workbook
        from build_v2_report import ExcelWriter

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "multi_sheet.xlsx"
            writer = ExcelWriter(output_path)

            writer.write_dataframe(pd.DataFrame({"a": [1]}), "Sheet1")
            writer.write_dataframe(pd.DataFrame({"b": [2]}), "Sheet2")
            writer.write_dataframe(pd.DataFrame({"c": [3]}), "Sheet3")
            writer.save()

            wb = load_workbook(output_path, read_only=True)
            assert len(wb.sheetnames) == 3
            assert "Sheet1" in wb.sheetnames
            assert "Sheet2" in wb.sheetnames
            assert "Sheet3" in wb.sheetnames
            wb.close()

    def test_writer_empty_dataframe(self):
        """ExcelWriter 写入空 DataFrame"""
        from build_v2_report import ExcelWriter

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "empty.xlsx"
            writer = ExcelWriter(output_path)

            writer.write_dataframe(pd.DataFrame(), "空Sheet")
            writer.save()

            assert output_path.exists()

    def test_writer_header_styling(self):
        """ExcelWriter：表头有样式（非空文件即验证）"""
        from openpyxl import load_workbook
        from build_v2_report import ExcelWriter

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "styled.xlsx"
            writer = ExcelWriter(output_path)

            df = pd.DataFrame({"列A": [1, 2], "列B": ["x", "y"]})
            writer.write_dataframe(df, "样式测试")
            writer.save()

            wb = load_workbook(output_path)
            ws = wb["样式测试"]

            # 表头行应该有填充色
            header_cell = ws.cell(row=1, column=1)
            assert header_cell.fill is not None
            # 表头字体加粗
            assert header_cell.font.bold

            wb.close()

    def test_writer_freeze_panes(self):
        """ExcelWriter：默认冻结首行"""
        from openpyxl import load_workbook
        from build_v2_report import ExcelWriter

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "freeze.xlsx"
            writer = ExcelWriter(output_path)

            df = pd.DataFrame({"a": [1], "b": [2]})
            writer.write_dataframe(df, "FreezeTest")
            writer.save()

            wb = load_workbook(output_path)
            ws = wb["FreezeTest"]
            assert ws.freeze_panes is not None
            wb.close()


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
