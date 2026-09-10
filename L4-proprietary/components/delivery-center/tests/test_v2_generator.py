"""
BDMS 交付月报 V2 生成器测试

验证 v2 生成器核心功能：
- 数据加载
- 状态引擎
- 映射引擎
- 考核引擎
- 报告生成
"""

import pytest
import pandas as pd
from pathlib import Path
import sys
import os

# 添加 v2 路径
V2_DIR = Path(__file__).parent.parent / "src" / "delivery_center" / "v2"
sys.path.insert(0, str(V2_DIR))
sys.path.insert(0, str(V2_DIR / "engines"))
sys.path.insert(0, str(V2_DIR / "utils"))
sys.path.insert(0, str(V2_DIR / "config"))


class TestStatusEngine:
    """状态引擎测试"""

    def test_determine_status_normal_delivery(self):
        """正常交付状态判定"""
        from status_engine import determine_status
        
        row = pd.Series({
            "状态": "交付邮件已归档",
            "履约项异常/变更类型": "",
            "预估交付完成日期": "2026-05-01",
            "预估验收完成日期": "2026-07-01",
            "实际服务/授权结束日期": "2026-12-31",
        })
        result = determine_status(row, "2026-06-30")
        assert result == "正常交付", f"预期 '正常交付', 实际 '{result}'"

    def test_determine_status_normal_acceptance(self):
        """正常验收状态判定"""
        from status_engine import determine_status
        
        row = pd.Series({
            "状态": "验收文件已归档",
            "履约项异常/变更类型": "",
            "预估交付完成日期": "2026-05-01",
            "预估验收完成日期": "2026-06-01",
            "实际服务/授权结束日期": "2026-12-31",
        })
        result = determine_status(row, "2026-06-30")
        # 验收已归档且服务期未结束 → 正常验收或正常服务
        assert result in ["正常验收", "正常服务"], f"预期正常验收/正常服务, 实际 '{result}'"

    def test_determine_status_delivery_abnormal(self):
        """交付异常状态判定"""
        from status_engine import determine_status
        
        row = pd.Series({
            "状态": "实施进行中",
            "履约项异常/变更类型": "履约项交付异常",
            "预估交付完成日期": "2026-05-01",
            "预估验收完成日期": "2026-07-01",
            "实际服务/授权结束日期": "2026-12-31",
        })
        result = determine_status(row, "2026-06-30")
        assert result == "交付异常", f"预期 '交付异常', 实际 '{result}'"

    def test_add_status_columns(self):
        """状态列添加测试"""
        from status_engine import add_status_columns
        
        df = pd.DataFrame({
            "状态": ["交付邮件已归档", "验收文件已归档", "实施进行中"],
            "履约项异常/变更类型": ["", "", "履约项交付异常"],
            "预估交付完成日期": ["2026-05-01", "2026-05-01", "2026-05-01"],
            "预估验收完成日期": ["2026-07-01", "2026-06-01", "2026-07-01"],
            "实际服务/授权结束日期": ["2026-12-31", "2026-12-31", "2026-12-31"],
            "项目状态": ["进行中", "进行中", "进行中"],
        })
        
        df = add_status_columns(df, "2026-06-30")
        
        # 验证关键列存在
        assert "履约项统计状态（即，财报-交付/确收状态）" in df.columns
        assert "1：正常交付" in df.columns
        assert "4：正常验收" in df.columns
        assert "3：交付异常" in df.columns
        assert "项目验收状态（即，财报-验收状态）" in df.columns
        
        # 验证状态列值带编号前缀
        status_val = df.iloc[0]["履约项统计状态（即，财报-交付/确收状态）"]
        assert status_val == "1：正常交付", f"预期带编号前缀, 实际 '{status_val}'"
        
        # 验证状态标记列（0/1）
        assert df.iloc[0]["1：正常交付"] == 1
        assert df.iloc[2]["3：交付异常"] == 1


class TestMappingEngine:
    """映射引擎测试"""

    def test_pm_dept_map_loaded(self):
        """项目经理-部门映射表加载"""
        from mapping_engine import load_pm_dept_map
        
        pm_dept = load_pm_dept_map()
        assert isinstance(pm_dept, dict)
        assert len(pm_dept) > 0, "映射表不应为空"

    def test_add_mapping_columns(self):
        """映射列添加"""
        from mapping_engine import add_mapping_columns, load_pm_dept_map
        
        pm_dept = load_pm_dept_map()
        # 找一个已知的 PM
        test_pm = list(pm_dept.keys())[0] if pm_dept else "测试"
        
        df = pd.DataFrame({
            "负责人": [test_pm, "未知人员"],
            "责任销售所属团队": ["北区营销部", "南区营销部"],
        })
        
        df = add_mapping_columns(df)
        
        assert "项目经理" in df.columns
        assert "项目经理所属部门" in df.columns
        assert "销售团队-统计" in df.columns
        
        # 验证已知 PM 的部门映射
        if pm_dept:
            assert df.iloc[0]["项目经理所属部门"] == pm_dept[test_pm]
        # 验证未知人员
        assert df.iloc[1]["项目经理所属部门"] == "未匹配"


class TestScoringEngine:
    """考核引擎测试"""

    def test_add_scoring_columns(self):
        """考核列添加"""
        from scoring_engine import add_scoring_columns
        
        df = pd.DataFrame({
            "基线-预估结项日期": ["2026-06-01", "2026-06-15", ""],
            "实际结项日期": ["2026-06-03", "2026-06-10", ""],
            "预估交付完成日期": ["2026-06-01", "2026-06-15", ""],
            "实际服务/授权开始日期": ["2026-06-05", "2026-06-12", ""],
        })
        
        df = add_scoring_columns(df)
        
        # 验证关键列存在
        assert "交付计划准确率“差异”" in df.columns
        assert "交付计划准确率“提前/延后”" in df.columns
        assert "交付计划准确率“是否跨月”" in df.columns
        assert "按时交付率“差异”" in df.columns
        assert "按时交付率“提前/延后”" in df.columns
        
        # 验证差异计算
        diff_val = df.iloc[0]["交付计划准确率“差异”"]
        assert pd.notna(diff_val), "差异值不应为空"


class TestSimpleComputedColumns:
    """简单计算列测试"""

    def test_add_simple_computed_columns(self):
        """简单计算列添加"""
        from mapping_engine import add_simple_computed_columns
        
        df = pd.DataFrame({
            "所属项目": ["SSXM-2026-0001", "SSXM-2026-0002"],
            "销售合同编号": ["HT-001", "HT-002"],
            "合同归档日期": ["2026-06-01", "2026-05-15"],
        })
        
        df = add_simple_computed_columns(df)
        
        assert "项目编号" in df.columns
        assert "统计项目编号" in df.columns
        assert "统计合同编号" in df.columns
        assert "合同归档年度" in df.columns
        
        # 验证合同归档年度
        assert df.iloc[0]["合同归档年度"] == "2026"
        assert df.iloc[1]["合同归档年度"] == "2026"


class TestExceptionLookup:
    """异常关联测试"""

    def test_add_exception_lookup_columns(self):
        """异常关联列添加"""
        from mapping_engine import add_exception_lookup_columns
        
        df_sign = pd.DataFrame({
            "销售合同编号": ["HT-001", "HT-002", "HT-003"],
        })
        
        df_exc = pd.DataFrame({
            "销售合同编号": ["HT-001", "HT-001"],  # HT-001 有异常
            "异常处置状态": ["处理中", "已完成"],
            "异常影响情况": ["进度延迟", "质量问题"],
            "交付说明": ["需延期", ""],
        })
        
        df = add_exception_lookup_columns(df_sign, df_exc)
        
        assert "异常项目对比" in df.columns
        assert "异常处置状态" in df.columns
        assert "异常影响情况" in df.columns
        
        # HT-001 应该有异常标记
        assert df.iloc[0]["异常项目对比"] == "有"
        # HT-002 应该没有
        assert df.iloc[1]["异常项目对比"] == ""


class TestReportGenerator:
    """报告生成器集成测试（轻量，用模拟数据）"""

    def test_generator_imports(self):
        """生成器模块可导入"""
        # 验证所有引擎可导入
        from engines.status_engine import add_status_columns
        from engines.scoring_engine import add_scoring_columns
        from engines.mapping_engine import add_mapping_columns
        from engines.exception_engine import build_exception_df
        from engines.handover_engine import build_revenue_handover_df, build_acceptance_handover_df
        
        assert True  # 不报错就是通过

    def test_sheet_count(self):
        """验证生成的报告有 15 个 Sheet（用真实数据跑一次）"""
        # 如果有真实数据就跑，没有就跳过
        ones_dir = Path.home() / ".openclaw" / "data" / "ones_exports"
        if not (ones_dir / "202606周报-签约项目统计.csv").exists():
            pytest.skip("缺少测试数据，跳过集成测试")
        
        from delivery_report_generator import generate_delivery_report
        from openpyxl import load_workbook
        
        output_path = generate_delivery_report("202606")
        assert output_path.exists()
        
        wb = load_workbook(output_path, read_only=True)
        assert len(wb.sheetnames) == 15, f"预期 15 个 Sheet, 实际 {len(wb.sheetnames)}"
        
        # 验证关键 Sheet 存在
        expected_sheets = [
            "签约", "POC&提前实施", "异常项目",
            "确收交接", "验收交接",
            "签约统计", "异常统计", "图例"
        ]
        for sheet in expected_sheets:
            assert sheet in wb.sheetnames, f"缺少 Sheet: {sheet}"
        
        wb.close()


if __name__ == "__main__":
    pytest.main([__file__, "-v"])


# ============================================================
# JSON 导出 + 多维明细导出 测试
# ============================================================
import json
import tempfile
from pathlib import Path

class TestJsonExport:
    """JSON 导出功能测试"""

    def test_export_report_json_basic(self):
        """基础 JSON 导出：sheet 数据正确序列化"""
        from generators.build_v2_report import export_report_json
        import pandas as pd

        df = pd.DataFrame({
            "合同编号": ["BC-001", "BC-002"],
            "金额": [10000, 20000],
            "日期": pd.to_datetime(["2026-01-01", "2026-02-01"]),
        })
        sheet_data = {"签约": df}

        with tempfile.TemporaryDirectory() as tmpdir:
            out = Path(tmpdir) / "report.json"
            result = export_report_json(sheet_data, out)

            assert result.exists()
            data = json.loads(result.read_text(encoding="utf-8"))

            assert "签约" in data
            assert data["签约"]["row_count"] == 2
            assert data["签约"]["col_count"] == 3
            assert data["签约"]["columns"] == ["合同编号", "金额", "日期"]
            assert len(data["签约"]["data"]) == 2
            assert data["签约"]["data"][0]["合同编号"] == "BC-001"
            assert "_meta" in data
            assert data["_meta"]["version"] == "v2"
            assert data["_meta"]["sheet_count"] == 1

    def test_export_report_json_skip_empty(self):
        """空 DataFrame 默认不导出"""
        from generators.build_v2_report import export_report_json
        import pandas as pd

        sheet_data = {
            "签约": pd.DataFrame({"a": [1]}),
            "空表": pd.DataFrame(),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            out = Path(tmpdir) / "report.json"
            export_report_json(sheet_data, out)
            data = json.loads(out.read_text(encoding="utf-8"))
            assert "签约" in data
            assert "空表" not in data

    def test_export_report_json_include_empty(self):
        """include_empty=True 时空表也导出"""
        from generators.build_v2_report import export_report_json
        import pandas as pd

        sheet_data = {"空表": pd.DataFrame(columns=["a", "b"])}
        with tempfile.TemporaryDirectory() as tmpdir:
            out = Path(tmpdir) / "report.json"
            export_report_json(sheet_data, out, include_empty=True)
            data = json.loads(out.read_text(encoding="utf-8"))
            assert "空表" in data
            assert data["空表"]["row_count"] == 0


class TestDetailExport:
    """多维明细导出测试"""

    def test_export_by_dimension_csv(self):
        """按维度拆分 CSV 导出"""
        from generators.build_v2_report import export_detail_by_dimension
        import pandas as pd

        df = pd.DataFrame({
            "项目名称": ["P1", "P2", "P3", "P4"],
            "项目经理所属部门": ["政企一部", "政企二部", "政企一部", "未分类"],
            "金额": [100, 200, 300, 400],
        })
        with tempfile.TemporaryDirectory() as tmpdir:
            out_dir = Path(tmpdir) / "details"
            files = export_detail_by_dimension(
                df, "项目经理所属部门", out_dir, format="csv", prefix="签约明细"
            )

            assert len(files) == 3
            assert out_dir.exists()
            # 政企一部有 2 个项目
            qiyi_file = out_dir / "签约明细_政企一部.csv"
            assert qiyi_file.exists()
            content = qiyi_file.read_text(encoding="utf-8-sig")
            assert "P1" in content
            assert "P3" in content

    def test_export_by_dimension_json(self):
        """按维度拆分 JSON 导出"""
        from generators.build_v2_report import export_detail_by_dimension
        import pandas as pd

        df = pd.DataFrame({
            "项目名称": ["P1", "P2"],
            "产品分类": ["安全", "数据"],
        })
        with tempfile.TemporaryDirectory() as tmpdir:
            out_dir = Path(tmpdir) / "details"
            files = export_detail_by_dimension(
                df, "产品分类", out_dir, format="json", prefix="产品"
            )
            assert len(files) == 2
            for f in files:
                assert f.suffix == ".json"
                data = json.loads(f.read_text(encoding="utf-8"))
                assert isinstance(data, list)
                assert len(data) == 1

    def test_export_by_dimension_empty_df(self):
        """空 DataFrame 返回空列表"""
        from generators.build_v2_report import export_detail_by_dimension
        import pandas as pd

        with tempfile.TemporaryDirectory() as tmpdir:
            files = export_detail_by_dimension(
                pd.DataFrame(), "部门", Path(tmpdir), format="csv"
            )
            assert files == []

    def test_export_by_dimension_invalid_format(self):
        """不支持的格式抛 ValueError"""
        from generators.build_v2_report import export_detail_by_dimension
        import pandas as pd

        df = pd.DataFrame({"部门": ["A"], "x": [1]})
        with tempfile.TemporaryDirectory() as tmpdir:
            import pytest
            with pytest.raises(ValueError, match="不支持的格式"):
                export_detail_by_dimension(df, "部门", Path(tmpdir), format="pdf")
