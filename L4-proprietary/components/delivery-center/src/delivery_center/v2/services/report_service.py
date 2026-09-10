"""
报告服务层

封装月报生成、任务管理、概要提取等业务逻辑。
异步任务用线程池执行，不阻塞 HTTP 请求。
"""

import sys
import json
import threading
import traceback
from pathlib import Path
from typing import Optional, List, Dict
from concurrent.futures import ThreadPoolExecutor

import pandas as pd

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

from db import (
    init_db, create_job, update_job_status, get_job, list_jobs,
    save_summaries, get_summaries_by_job,
)

# 报告输出目录（与 delivery_report_generator 一致）
REPORT_OUTPUT_DIR = Path.home() / ".openclaw" / "data" / "reports"

# 线程池（异步生成用）
_executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="bdms-report")
_init_done = False
_lock = threading.Lock()


def _ensure_init():
    """延迟初始化元数据库"""
    global _init_done
    if not _init_done:
        with _lock:
            if not _init_done:
                init_db()
                _init_done = True


def _generate_report_async(job_id: int, month: str):
    """异步生成报告（在线程池中执行）"""
    try:
        update_job_status(job_id, "running", progress=10)

        # 调用 v2 生成器
        sys.path.insert(0, str(BASE_DIR))
        from delivery_report_generator import generate_delivery_report

        update_job_status(job_id, "running", progress=30)
        output_path = generate_delivery_report(month)

        update_job_status(job_id, "running", progress=80)

        # 提取概要
        summaries = _extract_report_summaries(output_path, month)
        if summaries:
            save_summaries(job_id, month, summaries)

        update_job_status(job_id, "completed", progress=100, file_path=str(output_path))

    except Exception as e:
        error_msg = f"{e}\n{traceback.format_exc()}"
        update_job_status(job_id, "failed", error_msg=str(e))
        print(f"❌ 报告生成失败 job={job_id}: {e}")


def _extract_report_summaries(file_path: Path, month: str) -> List[Dict]:
    """从生成的 Excel 中提取各 Sheet 概要数据"""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
    except Exception:
        return []

    summaries = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        row_count = ws.max_row - 1 if ws.max_row > 0 else 0  # 减去表头

        # 提取关键指标（前几行前几列的数值型数据）
        summary = {}
        try:
            if row_count > 0 and ws.max_column >= 2:
                # 统计 Sheet：第一列是分类，第二列是数值
                numeric_values = []
                for row_idx in range(2, min(ws.max_row + 1, 20)):
                    cell_val = ws.cell(row=row_idx, column=2).value
                    if cell_val is not None and isinstance(cell_val, (int, float)):
                        numeric_values.append({
                            "label": ws.cell(row=row_idx, column=1).value or "",
                            "value": cell_val
                        })
                if numeric_values:
                    summary["top_metrics"] = numeric_values[:8]
        except Exception:
            pass

        summaries.append({
            "sheet_name": sheet_name,
            "row_count": row_count,
            "summary": summary,
        })

    wb.close()
    return summaries


# ============================================================
# 对外 API
# ============================================================

def generate_report_async(month: str) -> int:
    """触发生成报告，返回 job_id"""
    _ensure_init()
    job_id = create_job(month)
    _executor.submit(_generate_report_async, job_id, month)
    return job_id


def get_report_status(job_id: int) -> Optional[Dict]:
    """获取报告生成状态"""
    _ensure_init()
    job = get_job(job_id)
    if not job:
        return None
    return job


def list_reports(limit: int = 20) -> List[Dict]:
    """列出报告列表"""
    _ensure_init()
    return list_jobs(limit=limit)


def get_report_file_path(job_id: int) -> Optional[Path]:
    """获取报告文件路径（仅 completed 状态）"""
    job = get_report_status(job_id)
    if not job or job["status"] != "completed" or not job["file_path"]:
        return None
    path = Path(job["file_path"])
    return path if path.exists() else None


def get_report_summary(job_id: int) -> Optional[Dict]:
    """获取报告概要（用于预览页）"""
    _ensure_init()
    job = get_job(job_id)
    if not job:
        return None

    summaries = get_summaries_by_job(job_id)

    # 如果没有缓存的概要且报告已完成，尝试现场提取
    if not summaries and job["status"] == "completed" and job["file_path"]:
        file_path = Path(job["file_path"])
        if file_path.exists():
            summaries = _extract_report_summaries(file_path, job["month"])
            if summaries:
                save_summaries(job_id, job["month"], summaries)

    return {
        "job": job,
        "sheets": summaries,
    }


def get_report_download_name(job_id: int) -> str:
    """获取下载文件名"""
    job = get_report_status(job_id)
    if not job:
        return "report.xlsx"
    return f"交付月报-{job['month']}-v2.xlsx"
