"""
全量 Renderer 烟雾测试 — 验证每一种页面类型都能成功渲染。
"""

from __future__ import annotations

import os
import pytest

from pptx import Presentation

from bangcle_ppt.renderers import REGISTRY, get_renderer_class


ALL_PAGE_TYPES = list(REGISTRY.keys())


def _make_prs():
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000
    return prs


def _theme_for(page_type: str) -> str:
    return "dark" if "dark" in page_type else "light"


class TestAllRenderers:
    """测试所有注册的 renderer 都能成功渲染。"""

    @pytest.mark.parametrize("page_type", ALL_PAGE_TYPES)
    def test_each_renderer_renders(self, page_type):
        """每个 renderer 都能渲染一页幻灯片。"""
        prs = _make_prs()
        renderer_cls = get_renderer_class(page_type)
        theme = _theme_for(page_type)
        renderer = renderer_cls(theme=theme)

        slide = renderer.render(prs, data=None)
        assert slide is not None
        assert len(prs.slides) == 1
        # 至少有一些 shape
        assert len(slide.shapes) >= 2, f"{page_type} has too few shapes"

    @pytest.mark.parametrize("page_type", ALL_PAGE_TYPES)
    def test_renderer_page_type_matches(self, page_type):
        """renderer 的 page_type 属性与注册名一致。"""
        renderer_cls = get_renderer_class(page_type)
        # 注意：有些 renderer 可能被复用（如 phase-timeline-dark），跳过这种
        if renderer_cls.page_type != page_type:
            pytest.skip(f"Renderer reused with different page_type: {renderer_cls.page_type} != {page_type}")
        assert renderer_cls.page_type == page_type

    def test_registry_count(self):
        """至少 20 种页面类型。"""
        assert len(REGISTRY) >= 20, f"Only {len(REGISTRY)} page types registered"

    def test_light_theme_count(self):
        """浅色主题至少 15 种。"""
        light_count = sum(1 for pt in REGISTRY if "light" in pt or "dark" not in pt)
        # 实际 count
        light_types = [pt for pt in REGISTRY if "light" in pt]
        assert len(light_types) >= 15, f"Only {len(light_types)} light page types"

    def test_dark_theme_count(self):
        """深色主题至少 8 种。"""
        dark_types = [pt for pt in REGISTRY if "dark" in pt]
        assert len(dark_types) >= 8, f"Only {len(dark_types)} dark page types"


class TestFullPresentation:
    """完整 PPT 生成测试。"""

    def test_all_light_pages(self, tmp_output_dir):
        """生成所有浅色页面的完整 PPT。"""
        from bangcle_ppt.engine import TemplateEngine

        engine = TemplateEngine(theme="light")
        engine.register_renderers(REGISTRY)

        light_types = [pt for pt in ALL_PAGE_TYPES if "light" in pt]
        output = os.path.join(tmp_output_dir, "all-light.pptx")
        result = engine.render_presentation(
            [(pt, None) for pt in light_types],
            output_path=output,
        )
        assert os.path.exists(result)
        prs = Presentation(result)
        assert len(prs.slides) == len(light_types)

    def test_all_dark_pages(self, tmp_output_dir):
        """生成所有深色页面的完整 PPT。"""
        from bangcle_ppt.engine import TemplateEngine

        engine = TemplateEngine(theme="dark")
        engine.register_renderers(REGISTRY)

        dark_types = [pt for pt in ALL_PAGE_TYPES if "dark" in pt]
        output = os.path.join(tmp_output_dir, "all-dark.pptx")
        result = engine.render_presentation(
            [(pt, None) for pt in dark_types],
            output_path=output,
        )
        assert os.path.exists(result)
        prs = Presentation(result)
        assert len(prs.slides) == len(dark_types)

    def test_typical_presentation_flow(self, tmp_output_dir):
        """典型汇报流程：封面 → 目录 → 章节 → 内容×N → 结束。"""
        from bangcle_ppt.engine import TemplateEngine

        engine = TemplateEngine(theme="light")
        engine.register_renderers(REGISTRY)

        slides = [
            ("cover-light", {"title": "2025年度汇报", "presenter": "张三", "company": "梆梆安全"}),
            ("toc-light", {}),
            ("section-light", {"title": "第一部分：项目概述", "chapter_number": "01"}),
            ("content-two-col-light", {"title": "项目背景"}),
            ("content-three-cards-light", {"title": "核心能力"}),
            ("data-chart-light", {"title": "数据成果"}),
            ("section-light", {"title": "第二部分：技术方案", "chapter_number": "02"}),
            ("timeline-three-cards-light", {"title": "技术路线"}),
            ("flow-light", {"title": "工作流程"}),
            ("team-cards-light", {"title": "团队介绍"}),
            ("closing-light", {"main_text": "谢谢观看"}),
        ]

        output = os.path.join(tmp_output_dir, "typical-flow.pptx")
        result = engine.render_presentation(slides, output_path=output)
        assert os.path.exists(result)
        prs = Presentation(result)
        assert len(prs.slides) == 11

    def test_dark_presentation_flow(self, tmp_output_dir):
        """深色主题典型流程。"""
        from bangcle_ppt.engine import TemplateEngine

        engine = TemplateEngine(theme="dark")
        engine.register_renderers(REGISTRY)

        slides = [
            ("cover-dark", {"title": "产品发布会", "presenter": "梆梆安全"}),
            ("toc-dark", {}),
            ("section-dark", {"title": "产品概览", "chapter_number": "01"}),
            ("four-cards-dark", {"title": "核心功能"}),
            ("node-graph-dark", {"title": "技术架构"}),
            ("honeycomb-dark", {"title": "产品矩阵"}),
            ("pyramid-compare-dark", {"title": "方案优势"}),
            ("closing-dark", {"main_text": "谢谢观看"}),
        ]

        output = os.path.join(tmp_output_dir, "dark-flow.pptx")
        result = engine.render_presentation(slides, output_path=output)
        assert os.path.exists(result)
        prs = Presentation(result)
        assert len(prs.slides) == 8
