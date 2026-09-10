"""
模板引擎测试 — 验证渲染器能生成有效的 PPTX。
"""

from __future__ import annotations

import os
import pytest

from pptx import Presentation

from bangcle_ppt.engine import TemplateEngine
from bangcle_ppt.renderers import REGISTRY, get_renderer_class
from bangcle_ppt.theme import get_theme, ThemeVariant
from bangcle_ppt.base import RendererBase


class TestTheme:
    """主题系统测试。"""

    def test_get_light_theme(self):
        theme = get_theme("light")
        assert theme.variant == ThemeVariant.LIGHT
        assert theme.bg_color == "#FFFFFF"

    def test_get_dark_theme(self):
        theme = get_theme("dark")
        assert theme.variant == ThemeVariant.DARK
        assert theme.bg_color == "#00122B"

    def test_theme_colors_light(self):
        theme = get_theme("light")
        assert theme.title_color == "#2D74BB"
        assert theme.body_color == "#595757"
        assert theme.primary == "#2D74BB"

    def test_theme_colors_dark(self):
        theme = get_theme("dark")
        assert theme.title_color == "#FFFFFF"
        assert theme.body_color == "#FFFFFF"
        assert theme.accent_color == "#EFBA20"

    def test_hex_to_tuple(self):
        theme = get_theme("light")
        r, g, b = theme.hex_to_tuple("#2D74BB")
        assert r == 0x2D
        assert g == 0x74
        assert b == 0xBB


class TestRendererRegistry:
    """渲染器注册中心测试。"""

    def test_registry_not_empty(self):
        assert len(REGISTRY) >= 7  # 至少 7 种基础类型

    def test_core_types_exist(self):
        core_types = [
            "cover-light", "cover-dark",
            "toc-light", "toc-dark",
            "section-light", "section-dark",
            "content-two-col-light",
        ]
        for pt in core_types:
            assert pt in REGISTRY, f"Missing renderer: {pt}"

    def test_get_renderer_class(self):
        cls = get_renderer_class("cover-light")
        assert cls.page_type == "cover-light"

    def test_get_renderer_class_unknown(self):
        with pytest.raises(ValueError, match="Unknown page type"):
            get_renderer_class("nonexistent-type")

    def test_all_renderers_inherit_base(self):
        for pt, cls in REGISTRY.items():
            assert issubclass(cls, RendererBase), f"{pt} renderer does not inherit RendererBase"


class TestTemplateEngine:
    """模板引擎测试。"""

    @pytest.fixture
    def engine(self):
        engine = TemplateEngine(theme="light")
        engine.register_renderers(REGISTRY)
        return engine

    def test_engine_init(self, engine):
        assert engine.default_theme == "light"

    def test_register_renderer(self, engine):
        count_before = len(engine._renderers)
        assert count_before == len(REGISTRY)

    def test_get_renderer(self, engine):
        renderer = engine.get_renderer("cover-light")
        assert renderer is not None
        assert renderer.theme.variant == ThemeVariant.LIGHT

    def test_get_renderer_dark(self, engine):
        renderer = engine.get_renderer("cover-dark", theme="dark")
        assert renderer.theme.variant == ThemeVariant.DARK

    def test_get_renderer_unknown(self, engine):
        with pytest.raises(ValueError, match="No renderer registered"):
            engine.get_renderer("unknown")


class TestRenderersSmoke:
    """各渲染器烟雾测试：确保能生成有效幻灯片。"""

    def _make_prs(self):
        prs = Presentation()
        prs.slide_width = 12192000   # 13.333"
        prs.slide_height = 6858000   # 7.5"
        return prs

    @pytest.mark.parametrize("page_type", [
        "cover-light",
        "cover-dark",
        "toc-light",
        "toc-dark",
        "section-light",
        "section-dark",
        "content-two-col-light",
    ])
    def test_renderer_renders_slide(self, page_type):
        """每个注册的渲染器都能成功渲染一页。"""
        prs = self._make_prs()
        renderer_cls = get_renderer_class(page_type)
        theme_str = "dark" if "dark" in page_type else "light"
        renderer = renderer_cls(theme=theme_str)

        slide = renderer.render(prs)
        assert slide is not None
        # 确认幻灯片被添加到了 prs
        assert len(prs.slides) == 1

    def test_render_presentation_multi_slide(self, tmp_output_dir):
        """生成多页 PPT 并保存到文件。"""
        engine = TemplateEngine(theme="light")
        engine.register_renderers(REGISTRY)

        output_path = os.path.join(tmp_output_dir, "test_multi.pptx")
        result = engine.render_presentation(
            [
                ("cover-light", {"title": "测试封面"}),
                ("toc-light", {}),
                ("section-light", {"title": "第一章"}),
                ("content-two-col-light", {"title": "内容页"}),
            ],
            output_path=output_path,
        )

        assert os.path.exists(result)
        assert os.path.getsize(result) > 1000  # 至少有内容

        # 验证能打开
        prs = Presentation(result)
        assert len(prs.slides) == 4
        # 验证是 16:9
        assert prs.slide_width == 12192000
        assert prs.slide_height == 6858000

    def test_render_presentation_dark_theme(self, tmp_output_dir):
        """深色主题生成。"""
        engine = TemplateEngine(theme="dark")
        engine.register_renderers(REGISTRY)

        output_path = os.path.join(tmp_output_dir, "test_dark.pptx")
        result = engine.render_presentation(
            [
                ("cover-dark", {"title": "深色封面"}),
                ("toc-dark", {}),
                ("section-dark", {"title": "第一章"}),
            ],
            output_path=output_path,
        )

        assert os.path.exists(result)
        prs = Presentation(result)
        assert len(prs.slides) == 3

    def test_slide_has_shapes(self):
        """渲染后的幻灯片有 shape（不是空的）。"""
        prs = self._make_prs()
        renderer_cls = get_renderer_class("cover-light")
        renderer = renderer_cls(theme="light")
        slide = renderer.render(prs)

        # 封面页应该有多个 shape（至少标题、装饰线、logo 等）
        shape_count = len(slide.shapes)
        assert shape_count >= 3, f"Cover slide only has {shape_count} shapes"
