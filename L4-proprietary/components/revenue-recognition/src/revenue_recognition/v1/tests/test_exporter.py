"""Tests for revenue_recognition.v1.exporter module."""

import pytest
from pathlib import Path
from unittest.mock import patch, MagicMock
from openpyxl import Workbook
from revenue_recognition.v1.exporter import (
    RevenueExporter,
    HEADER_FONT,
    HEADER_FILL,
    DATA_FONT,
    THIN_BORDER,
    _apply_header_style,
    _apply_data_style,
)


class TestStyleConstants:
    """Test style constants are defined."""

    def test_header_font(self):
        assert HEADER_FONT.bold is True
        assert HEADER_FONT.size == 10

    def test_header_fill(self):
        assert HEADER_FILL.fill_type == "solid"

    def test_data_font(self):
        assert DATA_FONT.size == 10

    def test_thin_border(self):
        assert THIN_BORDER.left.style == "thin"
        assert THIN_BORDER.right.style == "thin"


class TestApplyStyles:
    """Test style helper functions."""

    def test_apply_header_style(self):
        wb = Workbook()
        ws = wb.active
        cell = ws.cell(row=1, column=1, value="test")
        _apply_header_style(cell)
        assert cell.font.bold is True
        assert cell.fill is not None
        assert cell.border.left.style == "thin"

    def test_apply_data_style(self):
        wb = Workbook()
        ws = wb.active
        cell = ws.cell(row=1, column=1, value="test")
        _apply_data_style(cell)
        assert cell.font.size == 10
        assert cell.border.left.style == "thin"


class TestRevenueExporter:
    """Test RevenueExporter class."""

    def test_default_init(self):
        exporter = RevenueExporter()
        assert exporter.engine is not None

    def test_custom_engine(self, engine_with_data):
        exporter = RevenueExporter(engine=engine_with_data)
        assert exporter.engine is engine_with_data

    @patch.object(RevenueExporter, '_build_legend')
    def test_export_creates_file(self, mock_legend, engine_with_data, tmp_path):
        output = tmp_path / "test_output.xlsx"
        exporter = RevenueExporter(engine=engine_with_data)
        result = exporter.export(period="202606", output_path=output)
        assert result.exists()
        assert result.suffix == ".xlsx"

    @patch.object(RevenueExporter, '_build_legend')
    def test_export_creates_sheets(self, mock_legend, engine_with_data, tmp_path):
        output = tmp_path / "test_output.xlsx"
        exporter = RevenueExporter(engine=engine_with_data)
        exporter.export(period="202606", output_path=output)

        from openpyxl import load_workbook
        wb = load_workbook(output)
        sheet_names = wb.sheetnames
        assert "汇总" in sheet_names
        assert "月度汇总记录" in sheet_names
        assert "履约汇总记录" in sheet_names
        assert "确收差异分析" in sheet_names
        assert "重拆履约" in sheet_names
        wb.close()

    @patch.object(RevenueExporter, '_build_legend')
    def test_export_summary_has_data(self, mock_legend, engine_with_data, tmp_path):
        output = tmp_path / "test_output.xlsx"
        exporter = RevenueExporter(engine=engine_with_data)
        exporter.export(period="202606", output_path=output)

        from openpyxl import load_workbook
        wb = load_workbook(output)
        ws = wb["汇总"]
        # Row 2 should have headers
        assert ws.cell(row=2, column=2).value == "期间"
        assert ws.cell(row=2, column=3).value == "新签合同"
        wb.close()

    @patch.object(RevenueExporter, '_build_legend')
    def test_export_default_output_path(self, mock_legend, engine_with_data, tmp_path):
        """Test that default output path is generated when not specified."""
        with patch("revenue_recognition.v1.exporter.OUTPUT_DIR", tmp_path):
            exporter = RevenueExporter(engine=engine_with_data)
            result = exporter.export(period="202606")
            assert result.exists()
            assert "202606" in result.name

    @patch.object(RevenueExporter, '_build_legend')
    def test_build_summary(self, mock_legend, engine_with_data):
        wb = Workbook()
        wb.remove(wb.active)
        exporter = RevenueExporter(engine=engine_with_data)
        exporter._build_summary(wb, "202606")
        ws = wb["汇总"]
        assert ws.cell(row=2, column=2).value == "期间"
        assert ws.cell(row=3, column=3).value == "新签合同额"

    def test_build_monthly_record(self, engine_with_data):
        wb = Workbook()
        wb.remove(wb.active)
        exporter = RevenueExporter(engine=engine_with_data)
        exporter._build_monthly_record(wb, "202606")
        ws = wb["月度汇总记录"]
        assert ws.cell(row=2, column=1).value == "统计期间"

    def test_build_performance_record(self, engine_with_data):
        wb = Workbook()
        wb.remove(wb.active)
        exporter = RevenueExporter(engine=engine_with_data)
        exporter._build_performance_record(wb, "202606")
        ws = wb["履约汇总记录"]
        assert ws.cell(row=1, column=1).value == "统计期间"

    def test_build_variance_analysis(self, engine_with_data):
        wb = Workbook()
        wb.remove(wb.active)
        exporter = RevenueExporter(engine=engine_with_data)
        exporter._build_variance_analysis(wb, "202606")
        ws = wb["确收差异分析"]
        assert ws.cell(row=1, column=1).value == "项目经理"

    def test_build_rebuild_perf(self, engine_with_data):
        wb = Workbook()
        wb.remove(wb.active)
        exporter = RevenueExporter(engine=engine_with_data)
        exporter._build_rebuild_perf(wb)
        ws = wb["重拆履约"]
        assert ws.cell(row=2, column=1).value == "合同编号"
