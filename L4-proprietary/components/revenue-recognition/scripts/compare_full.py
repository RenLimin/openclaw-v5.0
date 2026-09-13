#!/usr/bin/env python3
"""
全量对比脚本：自动化输出 vs 手工参考 Excel
支持 --period 202606 或 202605
"""

import argparse
import sys
import os
from pathlib import Path

# Add src to path
script_dir = Path(__file__).resolve().parent
project_dir = script_dir.parent
src_dir = project_dir / "src"
sys.path.insert(0, str(src_dir))

import openpyxl
from openpyxl.utils import get_column_letter

# ============================================================
# 参考文件路径
# ============================================================
REFERENCE_FILES = {
    "202606": Path(
        "/Users/bangcle/Bangcle Workspace/01. Management/2026/2026团队报告/202606/"
        "2026年计划确收&实际确收对比表202601-06-0724 - 差异分析.xlsx"
    ),
    "202605": Path(
        "/Users/bangcle/Bangcle Workspace/01. Management/2026/2026团队报告/202605/"
        "2026年计划确收&实际确收对比表202601-05-0627-差异分析.xlsx"
    ),
}

# ============================================================
# 对比配置
# ============================================================

# 汇总 sheet: 行范围与列范围
SUMMARY_CONFIG = {
    "data_rows": range(4, 18),    # Row 4-17 (月度+小计+合计)
    "data_cols": range(2, 13),    # B-L
    "skip_rows": {18, 19},        # 说明文字行
    "perf_rows": range(20, 29),   # Row 20-28 履约维度
    "perf_cols": range(3, 7),     # C-F
    "header_rows": {2, 3},
}


def is_numeric(val):
    return isinstance(val, (int, float))


def values_equal(auto_val, manual_val, tolerance=0.01):
    """比较两个值，数值用容差，字符串用 strip"""
    # None vs 0 视为相等
    if auto_val is None and manual_val is None:
        return True, 0.0
    if auto_val is None and manual_val == 0:
        return True, 0.0
    if manual_val is None and auto_val == 0:
        return True, 0.0
    if auto_val is None and manual_val is not None:
        return False, abs(float(manual_val)) if is_numeric(manual_val) else 0.0
    if manual_val is None and auto_val is not None:
        return False, abs(float(auto_val)) if is_numeric(auto_val) else 0.0

    # 数值比较
    if is_numeric(auto_val) and is_numeric(manual_val):
        diff = abs(float(auto_val) - float(manual_val))
        return diff <= tolerance, diff

    # 字符串比较
    a_str = str(auto_val).strip()
    m_str = str(manual_val).strip()
    return a_str == m_str, 0.0


def should_skip_cell(val):
    """跳过 #REF! 等无效单元格"""
    if isinstance(val, str) and "#REF" in val:
        return True
    return False


def compare_summary_sheet(ws_auto, ws_manual, tolerance=0.01):
    """对比汇总 sheet"""
    results = {
        "data": {"total": 0, "match": 0, "diff": 0, "diffs": []},
        "perf": {"total": 0, "match": 0, "diff": 0, "diffs": []},
    }

    # 数据区域 Row 4-17, Col B-L
    for row in SUMMARY_CONFIG["data_rows"]:
        for col in SUMMARY_CONFIG["data_cols"]:
            auto_val = ws_auto.cell(row=row, column=col).value
            manual_val = ws_manual.cell(row=row, column=col).value

            if auto_val is None and manual_val is None:
                continue
            if should_skip_cell(auto_val) or should_skip_cell(manual_val):
                continue

            results["data"]["total"] += 1
            is_match, diff = values_equal(auto_val, manual_val, tolerance)
            if is_match:
                results["data"]["match"] += 1
            else:
                results["data"]["diff"] += 1
                col_letter = get_column_letter(col)
                results["data"]["diffs"].append({
                    "cell": f"{col_letter}{row}",
                    "auto": auto_val,
                    "manual": manual_val,
                    "diff": diff,
                })

    # 履约维度 Row 20-28, Col C-F
    for row in SUMMARY_CONFIG["perf_rows"]:
        for col in SUMMARY_CONFIG["perf_cols"]:
            auto_val = ws_auto.cell(row=row, column=col).value
            manual_val = ws_manual.cell(row=row, column=col).value

            if auto_val is None and manual_val is None:
                continue
            if should_skip_cell(auto_val) or should_skip_cell(manual_val):
                continue

            # Row 20 是标题行 "履约维度："
            if row == 20 and col == 3:
                continue

            results["perf"]["total"] += 1
            is_match, diff = values_equal(auto_val, manual_val, tolerance)
            if is_match:
                results["perf"]["match"] += 1
            else:
                results["perf"]["diff"] += 1
                col_letter = get_column_letter(col)
                results["perf"]["diffs"].append({
                    "cell": f"{col_letter}{row}",
                    "auto": auto_val,
                    "manual": manual_val,
                    "diff": diff,
                })

    return results


def compare_row_by_row(ws_auto, ws_manual, sheet_name, 
                       start_row_auto, start_row_manual,
                       max_col, data_rows=None, tolerance=0.01):
    """逐行逐单元格对比（用于月度汇总记录、履约汇总记录等）"""
    results = {"total": 0, "match": 0, "diff": 0, "diffs": []}

    max_row_auto = ws_auto.max_row
    max_row_manual = ws_manual.max_row

    if data_rows:
        # 指定行范围
        rows_to_compare = data_rows
    else:
        # 取两个 sheet 的最大行数
        rows_to_compare = range(1, max(max_row_auto, max_row_manual) + 1)

    for row in rows_to_compare:
        for col in range(1, max_col + 1):
            auto_val = ws_auto.cell(row=row, column=col).value
            manual_val = ws_manual.cell(row=row, column=col).value

            if auto_val is None and manual_val is None:
                continue
            if should_skip_cell(auto_val) or should_skip_cell(manual_val):
                continue

            results["total"] += 1
            is_match, diff = values_equal(auto_val, manual_val, tolerance)
            if is_match:
                results["match"] += 1
            else:
                results["diff"] += 1
                col_letter = get_column_letter(col)
                results["diffs"].append({
                    "cell": f"{col_letter}{row}",
                    "auto": auto_val,
                    "manual": manual_val,
                    "diff": diff,
                })

    return results


def compare_sheets(wb_auto, wb_manual, period):
    """对比所有关键 sheet"""
    all_results = {}

    # ===== Sheet 1: 汇总 =====
    print("\n### 汇总 sheet")
    if "汇总" not in wb_auto.sheetnames:
        print("  ❌ 汇总 sheet 在自动化输出中不存在!")
        all_results["汇总"] = None
    else:
        ws_auto = wb_auto["汇总"]
        ws_manual = wb_manual["汇总"]
        result = compare_summary_sheet(ws_auto, ws_manual)
        total_all = result["data"]["total"] + result["perf"]["total"]
        match_all = result["data"]["match"] + result["perf"]["match"]
        diff_all = result["data"]["diff"] + result["perf"]["diff"]
        status = "PASS" if diff_all == 0 else "FAIL"
        print(f"  {status} | 总单元格: {total_all}, 匹配: {match_all}, 差异: {diff_all}")

        if result["data"]["diffs"]:
            print(f"  数据区域差异 ({len(result['data']['diffs'])} 项):")
            for d in result["data"]["diffs"][:20]:
                print(f"    {d['cell']}: 自动={d['auto']} vs 手工={d['manual']} (差={d['diff']:.6f})")

        if result["perf"]["diffs"]:
            print(f"  履约维度差异 ({len(result['perf']['diffs'])} 项):")
            for d in result["perf"]["diffs"][:10]:
                print(f"    {d['cell']}: 自动={d['auto']} vs 手工={d['manual']} (差={d['diff']:.6f})")

        all_results["汇总"] = result

    # ===== Sheet 2: 月度汇总记录 =====
    print("\n### 月度汇总记录 sheet")
    if "月度汇总记录" not in wb_auto.sheetnames:
        print("  ❌ 月度汇总记录 sheet 在自动化输出中不存在!")
        all_results["月度汇总记录"] = None
    else:
        ws_auto = wb_auto["月度汇总记录"]
        ws_manual = wb_manual["月度汇总记录"]
        # Row 1-2 是标题，数据从 Row 3 开始，Col A-K (1-11)
        result = compare_row_by_row(
            ws_auto, ws_manual, "月度汇总记录",
            start_row_auto=3, start_row_manual=3,
            max_col=11,
            data_rows=range(3, max(ws_auto.max_row, ws_manual.max_row) + 1),
        )
        status = "PASS" if result["diff"] == 0 else "FAIL"
        print(f"  {status} | 总单元格: {result['total']}, 匹配: {result['match']}, 差异: {result['diff']}")
        if result["diffs"]:
            print(f"  差异明细 (前 30 项):")
            for d in result["diffs"][:30]:
                print(f"    {d['cell']}: 自动={d['auto']} vs 手工={d['manual']} (差={d['diff']:.6f})")
        all_results["月度汇总记录"] = result

    # ===== Sheet 3: 履约汇总记录 =====
    print("\n### 履约汇总记录 sheet")
    if "履约汇总记录" not in wb_auto.sheetnames:
        print("  ❌ 履约汇总记录 sheet 在自动化输出中不存在!")
        all_results["履约汇总记录"] = None
    else:
        ws_auto = wb_auto["履约汇总记录"]
        ws_manual = wb_manual["履约汇总记录"]
        # Row 1 是标题，数据从 Row 2 开始，Col A-F (1-6)
        result = compare_row_by_row(
            ws_auto, ws_manual, "履约汇总记录",
            start_row_auto=2, start_row_manual=2,
            max_col=6,
            data_rows=range(2, max(ws_auto.max_row, ws_manual.max_row) + 1),
        )
        status = "PASS" if result["diff"] == 0 else "FAIL"
        print(f"  {status} | 总单元格: {result['total']}, 匹配: {result['match']}, 差异: {result['diff']}")
        if result["diffs"]:
            print(f"  差异明细 (前 30 项):")
            for d in result["diffs"][:30]:
                print(f"    {d['cell']}: 自动={d['auto']} vs 手工={d['manual']} (差={d['diff']:.6f})")
        all_results["履约汇总记录"] = result

    # ===== Sheet 4: 确收差异分析 =====
    print("\n### 确收差异分析 sheet")
    if "确收差异分析" not in wb_auto.sheetnames:
        print("  ❌ 确收差异分析 sheet 在自动化输出中不存在!")
        all_results["确收差异分析"] = None
    else:
        ws_auto = wb_auto["确收差异分析"]
        ws_manual = wb_manual["确收差异分析"]
        # Row 1-4 筛选条件, Row 5 空, Row 6-7 标题, Row 8+ 数据
        # 对比 Row 8+ 的数据区域
        data_start = 8
        result = compare_row_by_row(
            ws_auto, ws_manual, "确收差异分析",
            start_row_auto=data_start, start_row_manual=data_start,
            max_col=4,
            data_rows=range(data_start, max(ws_auto.max_row, ws_manual.max_row) + 1),
        )
        status = "PASS" if result["diff"] == 0 else "FAIL"
        print(f"  {status} | 总单元格: {result['total']}, 匹配: {result['match']}, 差异: {result['diff']}")
        if result["diffs"]:
            print(f"  差异明细:")
            for d in result["diffs"][:20]:
                print(f"    {d['cell']}: 自动={d['auto']} vs 手工={d['manual']} (差={d['diff']:.6f})")
        all_results["确收差异分析"] = result

    return all_results


def print_summary_report(results_202606, results_202605):
    """打印最终汇总报告"""
    print("\n" + "=" * 80)
    print("## 最终总结")
    print("=" * 80)

    for label, results in [("202606", results_202606), ("202605", results_202605)]:
        if results is None:
            print(f"\n### {label}: 未执行")
            continue

        total_diff = 0
        total_cells = 0
        for sheet_name, result in results.items():
            if result is None:
                continue
            if isinstance(result, dict) and "data" in result:
                # 汇总 sheet 的特殊结构
                total_diff += result["data"]["diff"] + result["perf"]["diff"]
                total_cells += result["data"]["total"] + result["perf"]["total"]
            elif isinstance(result, dict) and "diff" in result:
                total_diff += result["diff"]
                total_cells += result["total"]

        print(f"\n### {label}: 总差异 {total_diff} 项 (共 {total_cells} 单元格)")

        for sheet_name, result in results.items():
            if result is None:
                print(f"  {sheet_name}: MISSING")
                continue
            if isinstance(result, dict) and "data" in result:
                d = result["data"]["diff"] + result["perf"]["diff"]
                t = result["data"]["total"] + result["perf"]["total"]
                print(f"  {sheet_name}: {'PASS' if d == 0 else 'FAIL'} ({t} 单元格, {d} 差异)")
            elif isinstance(result, dict) and "diff" in result:
                print(f"  {sheet_name}: {'PASS' if result['diff'] == 0 else 'FAIL'} ({result['total']} 单元格, {result['diff']} 差异)")


def main():
    parser = argparse.ArgumentParser(description="全量对比：自动化输出 vs 手工参考 Excel")
    parser.add_argument("--period", required=True, choices=["202606", "202605"],
                        help="期间: 202606 或 202605")
    parser.add_argument("--auto-path", type=Path, default=None,
                        help="自动化输出路径（默认根据 period 推断）")
    parser.add_argument("--tolerance", type=float, default=0.01,
                        help="数值容差（默认 0.01）")
    args = parser.parse_args()

    period = args.period
    ref_path = REFERENCE_FILES[period]

    if args.auto_path:
        auto_path = args.auto_path
    else:
        auto_path = project_dir / "src" / "output" / f"确收自动化报表_{period}.xlsx"

    print("=" * 80)
    print(f"## {period} 端到端对比报告")
    print("=" * 80)
    print(f"自动化输出: {auto_path}")
    print(f"手工参考:   {ref_path}")
    print(f"容差:       {args.tolerance}")

    # 检查文件存在
    if not auto_path.exists():
        print(f"\n❌ 自动化输出不存在: {auto_path}")
        return 1

    if not ref_path.exists():
        print(f"\n❌ 手工参考文件不存在: {ref_path}")
        return 1

    # 加载工作簿
    print("\n加载自动化输出...")
    wb_auto = openpyxl.load_workbook(auto_path, data_only=True)
    print(f"  Sheets: {wb_auto.sheetnames}")

    print("加载手工参考...")
    wb_manual = openpyxl.load_workbook(ref_path, data_only=True)
    print(f"  Sheets: {wb_manual.sheetnames}")

    # 对比
    results = compare_sheets(wb_auto, wb_manual, period)

    wb_auto.close()
    wb_manual.close()

    # 汇总
    total_diff = 0
    total_cells = 0
    for sheet_name, result in results.items():
        if result is None:
            continue
        if isinstance(result, dict) and "data" in result:
            total_diff += result["data"]["diff"] + result["perf"]["diff"]
            total_cells += result["data"]["total"] + result["perf"]["total"]
        elif isinstance(result, dict) and "diff" in result:
            total_diff += result["diff"]
            total_cells += result["total"]

    print(f"\n{'=' * 80}")
    print(f"## {period} 总结")
    print(f"总单元格: {total_cells}, 总差异: {total_diff}")
    if total_diff == 0:
        print("✅ 全部通过！")
    else:
        print(f"❌ 存在 {total_diff} 项差异")
    print("=" * 80)

    return 0


if __name__ == "__main__":
    sys.exit(main())
