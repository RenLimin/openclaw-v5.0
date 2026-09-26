# 🔒 NO_TOKEN — 纯代码，零 AI 依赖
"""图例配置预落盘：从黄金基准 Sheet-15 提取 → md_reference。

对齐 DESIGN-DETAIL-DELIVERY-REPORT-v2.1.md §Step 1d。
用法: python3 -m bdms.modules.delivery_report.seed_legend [黄金基准.xlsx]
幂等：重复执行覆盖更新。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

from bdms.core import db as _db

# 默认黄金基准（202606）
DEFAULT_GOLDEN = Path(
    "/Users/bangcle/Bangcle Workspace/01. Management/2026/2026团队报告/"
    "202606/2026交付月报-20260630.xlsx"
)


def seed_legend(golden_path: Path = None, db_path: Path = None) -> dict:
    """从黄金基准提取图例配置 → md_reference（data_type='legend_config'）。"""
    golden_path = golden_path or DEFAULT_GOLDEN
    if not golden_path.exists():
        return {"error": f"黄金基准不存在: {golden_path}", "seeded": 0}

    wb = openpyxl.load_workbook(golden_path, read_only=True, data_only=True)
    if "图例" not in wb.sheetnames:
        wb.close()
        return {"error": "黄金基准无图例 Sheet", "seeded": 0}

    ws = wb["图例"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    conn = _db.get_connection(db_path)
    seeded = 0
    try:
        for r_idx, row in enumerate(rows):
            # 找非空行
            values = [str(v).strip() if v is not None else "" for v in row]
            if not any(values):
                continue
            code = values[0] if values[0] else None
            label = values[1] if len(values) > 1 and values[1] else None
            if not code and not label:
                continue

            # extra: 其余列作为附加属性
            extra = {}
            for i, v in enumerate(values[2:], start=2):
                if v:
                    extra[f"c{i}"] = v

            conn.execute(
                """INSERT INTO md_reference (data_type, code, label, extra, sort_order)
                   VALUES ('legend_config', ?, ?, ?, ?)
                   ON CONFLICT(data_type, code) DO UPDATE SET
                     label = excluded.label,
                     extra = excluded.extra,
                     sort_order = excluded.sort_order,
                     updated_at = datetime('now', 'localtime')""",
                (code or f"row_{r_idx}", label or "", 
                 json.dumps(extra, ensure_ascii=False) if extra else None,
                 r_idx),
            )
            seeded += 1
        conn.commit()
    finally:
        conn.close()

    return {"seeded": seeded, "source": str(golden_path)}


if __name__ == "__main__":
    golden = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    result = seed_legend(golden)
    print(json.dumps(result, ensure_ascii=False, indent=2))
