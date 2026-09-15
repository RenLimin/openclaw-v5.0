"""确认收入 —— 表头自适应导入器。

替代 revenue-recognition/v1/importer.py 的硬编码列号版本。
核心差异：按**列名**定位，不按列号；自动探测表头行。

数据落盘保持与原实现相同的表结构（plan_draft / budget_exec），
以便复用原实现的 engine.py 计算逻辑。
"""

import sys
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from bdms.core import db as _db
from bdms.core.header_mapper import detect_header_row, read_header, HeaderMapper


# ─── 目标字段 → 候选列名（按优先级）───
# 注意：候选名写成元组，index_any 会依次尝试

BUDGET_FIELDS: dict[str, tuple] = {
    "category":        ("分类",),
    "contract_no":     ("合同编号",),
    "contract_no_cal": ("合同编号(校准)", "合同编号（校准）"),
    "customer":        ("客户名称",),
    "end_user":        ("最终用户名称",),
    "sign_subject":    ("签约主体",),
    "archive_month":   ("合同归档月份",),
    "perf_id_budget":  ("履约ID(预算)", "履约ID（预算）"),
    "perf_detail_budget": ("履约明细(预算)", "履约明细（预算）"),
    "perf_id":         ("履约ID",),
    "rev_method":      ("收入确认方法",),
    "perf_amount":     ("单项履约义务金额",),
    "rev_prior":       ("截止20251231已确收金额",),
    "rev_future":      ("2026年及以后计划确收",),
    "no_plan":         ("年初-未立项&项目异常未计划确收",),
    "unrev_prior":     ("截止20251231未确收金额",),
    "unrev_adj":       ("截止20251231未确收金额(调整)",),
    "plan_start":      ("计划开始时间",),
    "plan_end":        ("计划结束时间",),
    "plan_done":       ("计划完成时间",),
    "year_est":        ("2026年预计",),
    "disappear_2026":  ("2026消失金额",),
    "disappear_future": ("2026年及以后消失金额",),
    "disappear_note":  ("消失备注",),
    "rebuild_perf":    ("重拆履约,提前和滞后同增", "重拆履约，提前和滞后同增"),
}

PLAN_FIELDS: dict[str, tuple] = {
    "note":            ("年初-填写说明",),
    "init_est_date":   ("年初-交付预计完成时间",),
    "est_date":        ("交付预计完成时间",),
    "contract_no":     ("合同编号",),
    "archive_month":   ("合同归档月份",),
    "contract_no2":    ("合同编号",),      # 第 2 个
    "prod_seq":        ("标准产品服务名称序号",),
    "perf_id":         ("履约ID",),
    "budget_perf_id":  ("对应预算履约ID",),
    "dept":            ("现行部门",),
    "contract_name":   ("合同名称",),
    "customer":        ("客户名称",),
    "end_user":        ("最终用户名称",),
    "contract_note":   ("合同备注",),
    "ops_note":        ("合同操作备注",),
    "tax_rate":        ("产品服务税率",),
    "sign_date":       ("合同签订日期",),
    "contract_start":  ("合同起始时间",),
    "contract_end":    ("合同结束时间",),
    "service_months":  ("服务期限(月)", "服务期限（月）"),
    "contract_type":   ("合同类型",),
    "version_type":    ("合同版本类型",),
    "gift":            ("是否赠送项项目",),
    "prod_category":   ("标准产品类别",),
    "prod_name":       ("合同产品服务名称",),
    "perf_detail":     ("履约义务明细",),
    "std_prod_name":   ("标准产品服务名称",),
    "rev_subject":     ("收入对应科目",),
    "tax_subject":     ("末级税金科目名称",),
    "price_basis":     ("价格拆分依据",),
    "accept_type":     ("验收文件类型",),
    "accept_term":     ("合同约定的验收条款",),
    "payment_term":    ("合同约定的收款节奏",),
    "rev_method":      ("收入确认方法",),
    "no_exec_reason":  ("履约不执行原因",),
    "qty_unit":        ("数量单位",),
    "qty":             ("数量",),
    "contract_amount": ("合同金额",),
    "confirm_amount":  ("确认合同额",),
    "perf_amount":     ("单项履约义务金额",),
    "plan_perf_amount": ("计划履约金额",),
    "rev_before_2025": ("截止20251231已确收",),
    "rev_2026_future": ("2026年及以后计划确收",),
    "plan_disappear":  ("计划-消失金额",),
    "disappear_reason": ("消失原因",),
}

# plan_draft 表字段顺序（与 db.py schema 对应）
PLAN_DB_COLUMNS = [
    "note", "init_est_date", "est_date", "contract_no", "archive_month",
    "contract_no2", "prod_seq", "perf_id", "budget_perf_id", "dept",
    "contract_name", "customer", "end_user", "contract_note", "ops_note",
    "tax_rate", "sign_date", "contract_start", "contract_end", "service_months",
    "contract_type", "version_type", "gift", "prod_category", "prod_name",
    "perf_detail", "std_prod_name", "rev_subject", "tax_subject", "price_basis",
    "accept_type", "accept_term", "payment_term", "rev_method", "no_exec_reason",
    "qty_unit", "qty", "contract_amount", "confirm_amount", "perf_amount",
    "plan_perf_amount", "rev_before_2025", "rev_2026_future",
    "plan_disappear", "disappear_reason",
]

BUDGET_DB_COLUMNS = [
    "category", "contract_no", "contract_no_cal", "customer", "end_user",
    "sign_subject", "archive_month", "perf_id_budget", "perf_detail_budget",
    "perf_id", "rev_method", "perf_amount", "rev_prior", "rev_future",
    "no_plan", "unrev_prior", "unrev_adj", "plan_start", "plan_end", "plan_done",
] + [f"m2026{m:02d}" for m in range(1, 13)] + [
    "year_est", "h1_plan", "h1_actual", "h1_ahead", "h1_behind",
] + [f"a2026{m:02d}" for m in range(1, 13)] + [
    "disappear_2026", "disappear_future", "disappear_note", "rebuild_perf",
    "forecast_category", "comparison_source",
]

# 数值型字段
NUMERIC_FIELDS = {
    "tax_rate", "service_months", "qty", "contract_amount", "confirm_amount",
    "perf_amount", "plan_perf_amount", "rev_before_2025", "rev_2026_future",
    "plan_disappear", "rev_prior", "rev_future", "no_plan", "unrev_prior",
    "unrev_adj", "year_est", "h1_plan", "h1_actual", "h1_ahead", "h1_behind",
    "disappear_2026", "disappear_future",
}


def _safe_float(val) -> Optional[float]:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).replace(",", "").strip())
    except (ValueError, TypeError):
        return None


def _safe_str(val) -> Optional[str]:
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


class AdaptiveImporter:
    """表头自适应的确收数据导入器。"""

    def __init__(self, db_path: Optional[Path] = None):
        self.db_path = Path(db_path) if db_path else None
        if self.db_path is None:
            from bdms.core.paths import DATA_DIR
            self.db_path = DATA_DIR / "revenue.db"

    # ─── 通用：读 sheet 为 (headers, rows) ───

    @staticmethod
    def _read_sheet(wb, sheet_name: str):
        """读取 sheet，自动探测表头行。返回 (mapper, headers, data_rows)。"""
        if sheet_name not in wb.sheetnames:
            return None, None, None
        ws = wb[sheet_name]
        hr = detect_header_row(ws)
        if hr is None:
            return None, None, None
        headers = read_header(ws, hr)
        mapper = HeaderMapper(headers)
        rows = []
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            rows.append(row)
        return mapper, headers, rows

    # ─── 导入预算执行表 ───

    def import_budget_exec(self, excel_path: Path, month: str) -> dict:
        wb = load_workbook(excel_path, read_only=True, data_only=True)
        mapper, headers, rows = self._read_sheet(wb, "预算执行表")
        if mapper is None:
            return {"rows": 0, "error": "找不到预算执行表或表头"}

        idx_map = {}   # db_col -> excel 索引(1-based)
        for field, candidates in BUDGET_FIELDS.items():
            idx = mapper.index_any(*candidates)
            if idx:
                idx_map[field] = idx

        # 月份列：第一个出现=计划月(m2026xx)，第二个=实际月(a2026xx)
        month_cols = mapper.find_month_columns("2026")

        conn = _db.get_connection(self.db_path)
        try:
            # 建表（复用原 schema）
            self._ensure_budget_table(conn)
            conn.execute("DELETE FROM budget_exec")
            n = 0
            for row in rows:
                rec = self._build_budget_record(row, idx_map, month_cols, month)
                if rec is None:
                    continue
                self._insert_budget(conn, rec)
                n += 1
            conn.commit()
            return {"rows": n, "columns": len(idx_map), "unmatched": sorted(mapper.unmatched)}
        finally:
            conn.close()

    def _build_budget_record(self, row, idx_map, month_cols, month):
        def get(field):
            i = idx_map.get(field)
            return row[i - 1] if i and i <= len(row) else None

        # 跳过空行/汇总行
        contract_no = _safe_str(get("contract_no"))
        if not contract_no or contract_no in ("合同编号", "合计", "总计"):
            return None

        rec = {}
        for field in BUDGET_FIELDS:
            val = get(field)
            rec[field] = _safe_float(val) if field in NUMERIC_FIELDS else _safe_str(val)

        # 月份列
        for m in range(1, 13):
            key = f"2026{m:02d}"
            rec[f"m2026{m:02d}"] = _safe_float(row[month_cols[f"{key}@1"] - 1]) \
                if f"{key}@1" in month_cols else None
            rec[f"a2026{m:02d}"] = _safe_float(row[month_cols[f"{key}@2"] - 1]) \
                if f"{key}@2" in month_cols and month_cols[f"{key}@2"] <= len(row) else None

        rec["comparison_source"] = month
        return rec

    def _ensure_budget_table(self, conn):
        """确保 budget_exec 表存在（schema 与原实现兼容）。"""
        conn.execute("""
            CREATE TABLE IF NOT EXISTS budget_exec (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                category TEXT, contract_no TEXT, contract_no_cal TEXT,
                customer TEXT, end_user TEXT, sign_subject TEXT,
                archive_month TEXT, perf_id_budget TEXT, perf_detail_budget TEXT,
                perf_id TEXT, rev_method TEXT, perf_amount REAL,
                rev_prior REAL, rev_future REAL, no_plan REAL,
                unrev_prior REAL, unrev_adj REAL,
                plan_start TEXT, plan_end TEXT, plan_done TEXT,
                m202601 REAL, m202602 REAL, m202603 REAL, m202604 REAL,
                m202605 REAL, m202606 REAL, m202607 REAL, m202608 REAL,
                m202609 REAL, m202610 REAL, m202611 REAL, m202612 REAL,
                year_est REAL, h1_plan REAL, h1_actual REAL, h1_ahead REAL, h1_behind REAL,
                a202601 REAL, a202602 REAL, a202603 REAL, a202604 REAL,
                a202605 REAL, a202606 REAL, a202607 REAL, a202608 REAL,
                a202609 REAL, a202610 REAL, a202611 REAL, a202612 REAL,
                disappear_2026 REAL, disappear_future REAL, disappear_note TEXT,
                rebuild_perf TEXT, forecast_category TEXT, comparison_source TEXT,
                imported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_budget_contract ON budget_exec(contract_no)")

    def _insert_budget(self, conn, rec):
        cols = [c for c in BUDGET_DB_COLUMNS if c in rec or c in ("comparison_source",)]
        placeholders = ",".join("?" for _ in cols)
        conn.execute(
            f"INSERT INTO budget_exec ({','.join(cols)}) VALUES ({placeholders})",
            [rec.get(c) for c in cols],
        )

    # ─── 导入计划确收底稿 ───

    def import_plan_draft(self, excel_path: Path, month: str) -> dict:
        wb = load_workbook(excel_path, read_only=True, data_only=True)
        mapper, headers, rows = self._read_sheet(wb, "计划确收底稿")
        if mapper is None:
            return {"rows": 0, "error": "找不到计划确收底稿或表头"}

        idx_map = {}
        for field, candidates in PLAN_FIELDS.items():
            # "合同编号" 第 2 次出现
            nth = 2 if field == "contract_no2" else 1
            idx = mapper.index_any(*candidates, nth=nth)
            if idx:
                idx_map[field] = idx

        conn = _db.get_connection(self.db_path)
        try:
            self._ensure_plan_table(conn)
            conn.execute("DELETE FROM plan_draft")
            n = 0
            for row in rows:
                rec = self._build_plan_record(row, idx_map)
                if rec is None:
                    continue
                self._insert_plan(conn, rec)
                n += 1
            conn.commit()
            return {"rows": n, "columns": len(idx_map), "unmatched": sorted(mapper.unmatched)}
        finally:
            conn.close()

    def _build_plan_record(self, row, idx_map):
        def get(field):
            i = idx_map.get(field)
            return row[i - 1] if i and i <= len(row) else None

        contract_no = _safe_str(get("contract_no"))
        if not contract_no or contract_no in ("合同编号", "合计", "总计"):
            return None

        rec = {}
        for field in PLAN_FIELDS:
            val = get(field)
            rec[field] = _safe_float(val) if field in NUMERIC_FIELDS else _safe_str(val)
        return rec

    def _ensure_plan_table(self, conn):
        cols_sql = ", ".join(f"{c} TEXT" for c in PLAN_DB_COLUMNS
                             if c not in ("tax_rate", "service_months", "qty",
                                          "contract_amount", "confirm_amount",
                                          "perf_amount", "plan_perf_amount",
                                          "rev_before_2025", "rev_2026_future",
                                          "plan_disappear"))
        numeric_sql = ", ".join(
            f"{c} REAL" for c in ["tax_rate", "service_months", "qty",
                                  "contract_amount", "confirm_amount",
                                  "perf_amount", "plan_perf_amount",
                                  "rev_before_2025", "rev_2026_future",
                                  "plan_disappear"]
        )
        conn.execute(f"""
            CREATE TABLE IF NOT EXISTS plan_draft (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                {cols_sql}, {numeric_sql},
                imported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_plan_contract ON plan_draft(contract_no)")

    def _insert_plan(self, conn, rec):
        cols = [c for c in PLAN_DB_COLUMNS if c in rec]
        placeholders = ",".join("?" for _ in cols)
        conn.execute(
            f"INSERT INTO plan_draft ({','.join(cols)}) VALUES ({placeholders})",
            [rec.get(c) for c in cols],
        )
