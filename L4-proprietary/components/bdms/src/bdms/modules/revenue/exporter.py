"""确认收入导出 —— RevenueReportExporter。"""

from pathlib import Path
from typing import Optional

from bdms.modules.revenue.summary_engine import SummaryEngine


class RevenueSummaryError(ValueError):
    """汇总数据不可用（无数据/未导入）时抛出，对应原 api.py 的 HTTP 400。"""


class RevenueReportExporter:
    """确收汇总Excel导出器。"""

    def __init__(self, db_path: Optional[Path] = None):
        pass

    def export(self, month: str, out_path: Optional[Path] = None) -> Path:
        """导出确收汇总为Excel，返回输出路径。"""
        import pandas as pd
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        r = SummaryEngine().compute_summary(month)
        if "error" in r:
            raise RevenueSummaryError(r["error"])

        wb = Workbook()

        # Sheet1: 汇总表
        ws = wb.active
        ws.title = "汇总表"

        rows_data = r.get("rows", [])
        if rows_data:
            # 表头
            headers = list(rows_data[0].keys())
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EDF5")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # 数据
            for row_idx, row in enumerate(rows_data, 2):
                for col_idx, key in enumerate(headers, 1):
                    ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))

            # 列宽自适应
            for col in ws.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 30)

        # Sheet2: 月度趋势
        ws2 = wb.create_sheet("月度趋势")
        monthly = r.get("monthly", {})
        months = monthly.get("months", [])
        if months:
            ws2.cell(row=1, column=1, value="月份").font = Font(bold=True)
            for i, cat in enumerate(["新签", "递延", "合计"], 2):
                ws2.cell(row=1, column=i*2-1, value=f"{cat}-计划").font = Font(bold=True)
                ws2.cell(row=1, column=i*2, value=f"{cat}-实际").font = Font(bold=True)
            for row_idx, m in enumerate(months, 2):
                ws2.cell(row_idx, column=1, value=m)
                for i, cat in enumerate(["新签", "递延", "合计"], 2):
                    series = monthly.get(cat, {})
                    plan_val = series.get("plan", [])[row_idx-2] if row_idx-2 < len(series.get("plan", [])) else ""
                    actual_val = series.get("actual", [])[row_idx-2] if row_idx-2 < len(series.get("actual", [])) else ""
                    ws2.cell(row=row_idx, column=i*2-1, value=plan_val)
                    ws2.cell(row=row_idx, column=i*2, value=actual_val)

        # 确定输出路径
        if out_path is None:
            from bdms.core import paths
            mdir = paths.month_dir(month)
            out_path = mdir / f"确收汇总_{month}.xlsx"

        wb.save(out_path)
        return out_path

