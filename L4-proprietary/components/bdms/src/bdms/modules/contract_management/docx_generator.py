"""合同 docx 生成器 — ContractDocxGenerator。

🔒 NO_TOKEN — 纯代码，零 AI 依赖。

功能：
  - 基于模板 + 占位符替换生成合同 Word 文档
  - 支持水印（草稿/已审批/已签署）
  - 使用 python-docx
"""

import os
from pathlib import Path
from typing import Optional
from datetime import datetime

from bdms.core.paths import OUTPUT_DIR, output_path


class ContractDocxGenerator:
    """合同 Word 文档生成器。"""

    # 占位符映射：{placeholder} → contract 字段名
    PLACEHOLDERS = {
        "{{contract_no}}": "contract_no",
        "{{title}}": "title",
        "{{party_a}}": "party_a",
        "{{party_b}}": "party_b",
        "{{amount}}": "_amount_formatted",
        "{{currency}}": "currency",
        "{{effective_date}}": "effective_date",
        "{{expiry_date}}": "expiry_date",
        "{{signed_date}}": "signed_date",
        "{{today}}": "_today",
        "{{contract_type}}": "contract_type",
    }

    def generate(self, contract: dict, template_path: str = None,
                 output_path: str = None) -> str:
        """生成合同 Word 文档。

        Args:
            contract: 合同数据 dict
            template_path: 模板文件路径（可选，无模板时使用默认结构）
            output_path: 输出路径（可选）

        Returns:
            生成的文件路径
        """
        try:
            from docx import Document
            from docx.shared import Pt, Inches, Cm, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            raise ImportError("请安装 python-docx: pip install python-docx")

        # 确定水印
        status = contract.get("status", "draft")
        watermark = self._get_watermark(status)

        if template_path and Path(template_path).exists():
            doc = Document(template_path)
            self._replace_placeholders(doc, contract)
        else:
            doc = self._build_default_doc(contract)

        # 添加水印
        if watermark:
            self._add_watermark(doc, watermark)

        # 保存
        if output_path:
            target = Path(output_path)
        else:
            contract_no = contract.get("contract_no", "unknown")
            target = OUTPUT_DIR / f"合同_{contract_no}_{status}.docx"

        target.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(target))
        return str(target)

    def _replace_placeholders(self, doc, contract: dict) -> None:
        """替换文档中的占位符。"""
        # 准备替换值
        values = self._prepare_values(contract)

        # 替换段落中的占位符
        for para in doc.paragraphs:
            for placeholder, field in self.PLACEHOLDERS.items():
                if placeholder in para.text:
                    value = values.get(field, "")
                    for run in para.runs:
                        if placeholder in run.text:
                            run.text = run.text.replace(placeholder, str(value))

        # 替换表格中的占位符
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for placeholder, field in self.PLACEHOLDERS.items():
                        if placeholder in cell.text:
                            cell.text = cell.text.replace(
                                placeholder, str(values.get(field, ""))
                            )

    def _build_default_doc(self, contract: dict):
        """构建默认合同文档结构（无模板时）。"""
        from docx import Document
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # 设置默认字体
        style = doc.styles["Normal"]
        font = style.font
        font.name = "微软雅黑"
        font.size = Pt(11)

        # 标题
        title = doc.add_heading(contract.get("title", "合同"), level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 合同编号
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(f"合同编号：{contract.get('contract_no', '')}")
        run.font.size = Pt(10)

        # 甲乙双方
        doc.add_paragraph("")
        p = doc.add_paragraph()
        run = p.add_run("甲方：")
        run.bold = True
        p.add_run(contract.get("party_a", "________"))

        p = doc.add_paragraph()
        run = p.add_run("乙方：")
        run.bold = True
        p.add_run(contract.get("party_b", "________"))

        # 鉴于条款
        doc.add_paragraph("")
        p = doc.add_paragraph()
        run = p.add_run("鉴于条款")
        run.bold = True
        run.font.size = Pt(12)
        doc.add_paragraph("甲乙双方本着平等互利、诚实信用的原则，经友好协商，就以下事项达成一致，特订立本合同。")

        # 主要条款
        doc.add_paragraph("")
        p = doc.add_paragraph()
        run = p.add_run("第一条 合同标的")
        run.bold = True
        doc.add_paragraph("（请在此处填写合同标的具体内容）")

        doc.add_paragraph("")
        p = doc.add_paragraph()
        run = p.add_run("第二条 合同金额")
        run.bold = True
        amount = contract.get("amount", 0)
        currency = contract.get("currency", "CNY")
        doc.add_paragraph(f"合同总金额为 {currency} {amount:,.2f} 元。")

        doc.add_paragraph("")
        p = doc.add_paragraph()
        run = p.add_run("第三条 合同期限")
        run.bold = True
        effective = contract.get("effective_date", "____年__月__日")
        expiry = contract.get("expiry_date", "____年__月__日")
        doc.add_paragraph(f"合同有效期自 {effective} 至 {expiry}。")

        # 签署区
        doc.add_paragraph("")
        doc.add_paragraph("")
        p = doc.add_paragraph()
        run = p.add_run("签署页")
        run.bold = True
        run.font.size = Pt(12)

        doc.add_paragraph("")
        doc.add_paragraph("甲方（盖章）：___________        乙方（盖章）：___________")
        doc.add_paragraph("")
        doc.add_paragraph("授权代表：___________          授权代表：___________")
        doc.add_paragraph("")
        signed_date = contract.get("signed_date") or datetime.now().strftime("%Y年%m月%d日")
        doc.add_paragraph(f"日期：{signed_date}                    日期：{signed_date}")

        return doc

    def _add_watermark(self, doc, text: str) -> None:
        """添加文字水印（页眉处）。"""
        try:
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
        except ImportError:
            return

        for section in doc.sections:
            header = section.header
            header.is_linked_to_previous = False
            p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(40)
            run.font.color.rgb = RGBColor(200, 200, 200)
            # 设置字体
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    @staticmethod
    def _get_watermark(status: str) -> str:
        """根据状态返回水印文字。"""
        return {
            "draft": "草 稿",
            "review1": "审批中",
            "review2": "审批中",
            "review3": "审批中",
            "approved": "已审批",
            "signed": "已签署",
            "archived": "已归档",
            "rejected": "已驳回",
        }.get(status, "")

    @staticmethod
    def _prepare_values(contract: dict) -> dict:
        """准备占位符替换值。"""
        values = dict(contract)
        amount = contract.get("amount", 0)
        values["_amount_formatted"] = f"{amount:,.2f}"
        values["_today"] = datetime.now().strftime("%Y年%m月%d日")
        return values
