"""合同管理 Excel 导出器 — ContractExporter。

🔒 NO_TOKEN — 纯代码，零 AI 依赖。

导出内容：
  - export_contract_overview — 合同概览 Excel
  - export_risk_details — 风险明细
  - export_approval_log — 审批日志
  - export_audit_trail — 审计追踪
"""

from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from bdms.core import db as _db
from bdms.core.paths import OUTPUT_DIR, output_path

from ..base import BaseExporter
from ._crypto import decrypt, mask_name, mask_amount


# ─── 样式常量 ───

HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF4472C4")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_FONT = Font(name="微软雅黑", size=10)
DATA_ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
DATA_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
DATA_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
THIN = Side(style="thin", color="FFD9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM_FMT_MONEY = '#,##0.00'

# 风险等级颜色
RISK_HIGH_FILL = PatternFill("solid", fgColor="FFFF0000")
RISK_MED_FILL = PatternFill("solid", fgColor="FFFFA500")
RISK_LOW_FILL = PatternFill("solid", fgColor="FF90EE90")


class ContractExporter(BaseExporter):
    """合同管理 Excel 导出器。"""

    module_name = "contract_management"
    sheet_order = ["合同概览", "风险明细", "审批日志", "审计追踪"]

    def __init__(self, db_path: Optional[Path] = None):
        super().__init__(db_path)

    # ─── 公开导出方法 ───

    def export_contract_overview(self, month: str = None,
                                  out_path: Optional[Path] = None) -> Path:
        """导出合同概览 Excel。"""
        wb = Workbook()
        wb.remove(wb.active)

        conn = _db.get_connection(self.db_path)
        try:
            self._write_overview_sheet(wb, conn, month)
        finally:
            conn.close()

        target = Path(out_path) if out_path else output_path(
            f"合同概览_{month or 'all'}.xlsx"
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
        return target

    def export_risk_details(self, contract_id: int = None,
                            out_path: Optional[Path] = None) -> Path:
        """导出风险明细 Excel。"""
        wb = Workbook()
        wb.remove(wb.active)

        conn = _db.get_connection(self.db_path)
        try:
            self._write_risk_sheet(wb, conn, contract_id)
        finally:
            conn.close()

        target = Path(out_path) if out_path else output_path(
            f"风险明细_{contract_id or 'all'}.xlsx"
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
        return target

    def export_approval_log(self, contract_id: int = None,
                            out_path: Optional[Path] = None) -> Path:
        """导出审批日志 Excel。"""
        wb = Workbook()
        wb.remove(wb.active)

        conn = _db.get_connection(self.db_path)
        try:
            self._write_approval_sheet(wb, conn, contract_id)
        finally:
            conn.close()

        target = Path(out_path) if out_path else output_path(
            f"审批日志_{contract_id or 'all'}.xlsx"
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
        return target

    def export_audit_trail(self, contract_id: int = None,
                           out_path: Optional[Path] = None) -> Path:
        """导出审计追踪 Excel。"""
        wb = Workbook()
        wb.remove(wb.active)

        conn = _db.get_connection(self.db_path)
        try:
            self._write_audit_sheet(wb, conn, contract_id)
        finally:
            conn.close()

        target = Path(out_path) if out_path else output_path(
            f"审计追踪_{contract_id or 'all'}.xlsx"
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
        return target

    # ─── BaseExporter 抽象方法 ───

    def _write_all_sheets(self, wb: Workbook, month: str, conn) -> None:
        """写所有 sheet（默认导出合同概览 + 风险明细）。"""
        self._write_overview_sheet(wb, conn, month)
        self._write_risk_sheet(wb, conn, contract_id=None)

    # ─── 合同概览 Sheet ───

    def _write_overview_sheet(self, wb: Workbook, conn, month: str = None) -> None:
        ws = wb.create_sheet("合同概览")

        headers = ["合同编号", "标题", "类型", "甲方", "乙方",
                    "金额", "币种", "状态", "审批级别", "生效日期", "到期日期", "签署日期"]

        # 写表头
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER

        # 查询数据
        where = "(deleted_at IS NULL OR deleted_at = '')"
        params: list = []
        if month:
            where += " AND substr(created_at, 1, 6) = ?"
            params.append(month)

        rows = conn.execute(
            f"""SELECT contract_no, title, contract_type,
                       party_a, party_b, amount, currency, status,
                       approval_level, effective_date, expiry_date, signed_date
                FROM cr_contracts
                WHERE {where}
                ORDER BY id DESC""",
            params,
        ).fetchall()

        status_map = {
            "draft": "起草", "review1": "一级审批", "review2": "二级审批",
            "review3": "三级审批", "approved": "审批通过", "signed": "已签署",
            "archived": "已归档", "rejected": "已驳回",
        }

        for ri, r in enumerate(rows, 2):
            values = [
                r["contract_no"],
                r["title"],
                r["contract_type"] or "",
                mask_name(decrypt(r["party_a"] or "")),
                mask_name(decrypt(r["party_b"] or "")),
                mask_amount(r["amount"] or 0),
                r["currency"] or "CNY",
                status_map.get(r["status"], r["status"]),
                r["approval_level"] or 1,
                r["effective_date"] or "",
                r["expiry_date"] or "",
                r["signed_date"] or "",
            ]
            for ci, val in enumerate(values, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.font = DATA_FONT
                cell.border = BORDER
                if ci >= 6 and ci <= 7:
                    cell.alignment = DATA_ALIGN_RIGHT
                else:
                    cell.alignment = DATA_ALIGN_LEFT

        ws.freeze_panes = "A2"
        col_widths = [18, 24, 12, 12, 12, 12, 6, 10, 8, 12, 12, 12]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ─── 风险明细 Sheet ───

    def _write_risk_sheet(self, wb: Workbook, conn, contract_id: int = None) -> None:
        ws = wb.create_sheet("风险明细")

        headers = ["扫描批次", "合同ID", "风险类别", "风险等级", "问题摘要", "建议", "状态", "扫描时间"]

        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER

        where = "(deleted_at IS NULL OR deleted_at = '')"
        params: list = []
        if contract_id:
            where += " AND contract_id = ?"
            params.append(contract_id)

        rows = conn.execute(
            f"""SELECT scan_batch, contract_id, risk_category, risk_level,
                       issue_summary, suggestion, status, created_at
                FROM cr_risk_scan_results
                WHERE {where}
                ORDER BY created_at DESC""",
            params,
        ).fetchall()

        for ri, r in enumerate(rows, 2):
            values = [
                r["scan_batch"], r["contract_id"], r["risk_category"],
                r["risk_level"], r["issue_summary"], r["suggestion"],
                r["status"], r["created_at"],
            ]
            for ci, val in enumerate(values, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.font = DATA_FONT
                cell.border = BORDER
                cell.alignment = DATA_ALIGN_LEFT

                # 风险等级着色
                if ci == 4:
                    level = r["risk_level"]
                    if level == "high":
                        cell.fill = RISK_HIGH_FILL
                    elif level == "medium":
                        cell.fill = RISK_MED_FILL
                    elif level == "low":
                        cell.fill = RISK_LOW_FILL

        ws.freeze_panes = "A2"
        col_widths = [18, 10, 14, 10, 30, 30, 10, 20]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ─── 审批日志 Sheet ───

    def _write_approval_sheet(self, wb: Workbook, conn, contract_id: int = None) -> None:
        ws = wb.create_sheet("审批日志")

        headers = ["合同ID", "原状态", "目标状态", "操作", "审批人", "审批角色", "意见", "操作时间"]

        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER

        where = "1=1"
        params: list = []
        if contract_id:
            where += " AND contract_id = ?"
            params.append(contract_id)

        rows = conn.execute(
            f"""SELECT contract_id, from_status, to_status, action,
                       approver_name, approver_role, comment, created_at
                FROM cr_approval_log
                WHERE {where}
                ORDER BY created_at DESC""",
            params,
        ).fetchall()

        for ri, r in enumerate(rows, 2):
            values = [
                r["contract_id"], r["from_status"], r["to_status"],
                r["action"], r["approver_name"], r["approver_role"] or "",
                r["comment"] or "", r["created_at"],
            ]
            for ci, val in enumerate(values, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.font = DATA_FONT
                cell.border = BORDER
                cell.alignment = DATA_ALIGN_LEFT

        ws.freeze_panes = "A2"
        col_widths = [10, 12, 12, 10, 12, 12, 30, 20]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ─── 审计追踪 Sheet ───

    def _write_audit_sheet(self, wb: Workbook, conn, contract_id: int = None) -> None:
        ws = wb.create_sheet("审计追踪")

        headers = ["合同ID", "操作", "字段名", "旧值", "新值", "操作人", "操作时间"]

        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER

        where = "1=1"
        params: list = []
        if contract_id:
            where += " AND contract_id = ?"
            params.append(contract_id)

        rows = conn.execute(
            f"""SELECT contract_id, operation, field_name, old_value,
                       new_value, operator, created_at
                FROM cr_audit_trail
                WHERE {where}
                ORDER BY created_at DESC""",
            params,
        ).fetchall()

        for ri, r in enumerate(rows, 2):
            values = [
                r["contract_id"], r["operation"], r["field_name"] or "",
                r["old_value"] or "", r["new_value"] or "",
                r["operator"], r["created_at"],
            ]
            for ci, val in enumerate(values, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.font = DATA_FONT
                cell.border = BORDER
                cell.alignment = DATA_ALIGN_LEFT

        ws.freeze_panes = "A2"
        col_widths = [10, 12, 14, 20, 20, 12, 20]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
