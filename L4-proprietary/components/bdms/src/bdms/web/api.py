"""BDMS Web API 路由。

按模块分组：
  /api/report/*       交付月报
  /api/revenue/*      确认收入
  /api/master-data/*  基础数据
  /api/dashboard/*    统计看板
  /api/settings/*     系统设定
"""

from typing import Any, Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel

from bdms.core import db as _db
from bdms.core.db import normalize_month

router = APIRouter(prefix="/api")


# ========== 请求模型 ==========

class GenerateRequest(BaseModel):
    month: str
    mode: str = "auto"


class SettingUpdate(BaseModel):
    key: str
    value: Any


class DriftQuery(BaseModel):
    month: str
    dimension: str
    value: str
    sheet: str = "签约"


# ========== 1. 交付月报 ==========

@router.get("/report/months")
async def report_months():
    from bdms.modules.delivery_report.service import DeliveryReportService
    return {"months": DeliveryReportService().list_months()}


@router.post("/report/generate")
async def report_generate(req: GenerateRequest):
    from bdms.modules.delivery_report.service import DeliveryReportService
    try:
        return DeliveryReportService().generate(normalize_month(req.month), req.mode)
    except Exception as e:
        raise HTTPException(500, str(e))


@router.get("/report/export/{month}")
async def report_export(month: str):
    month = normalize_month(month)
    from bdms.modules.delivery_report.exporter import DeliveryReportExporter
    try:
        path = DeliveryReportExporter().export(month)
    except Exception as e:
        raise HTTPException(500, str(e))
    return FileResponse(path, filename=f"交付月报_{month}.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get("/report/export-sheet/{month}/{sheet}")
async def report_export_sheet(month: str, sheet: str, search: str = ""):
    """导出单个 Sheet 为 CSV（支持搜索过滤）。"""
    import io
    import csv
    month = normalize_month(month)
    from bdms.modules.delivery_report.engine import DeliveryReportEngine
    engine = DeliveryReportEngine()
    data = engine.load(month)
    if sheet not in data:
        raise HTTPException(404, f"{month} 无 {sheet} 数据")
    df = data[sheet]
    
    # 搜索过滤
    if search and search.strip():
        keyword = search.strip()
        mask = df.apply(
            lambda row: any(keyword.lower() in str(v).lower() for v in row.values),
            axis=1
        )
        df = df[mask]
    
    # 导出 CSV
    output = io.StringIO()
    df.fillna("").to_csv(output, index=False, quoting=csv.QUOTE_ALL)
    content = output.getvalue().encode('utf-8-sig')
    
    from fastapi.responses import Response
    filename = f"{month}_{sheet}{'_筛选' if search else ''}.csv"
    return Response(
        content=content,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=\"{filename}\""}
    )


@router.get("/report/data/{month}/{sheet}")
async def report_data(month: str, sheet: str, page: int = 1, page_size: int = 50, search: str = ""):
    """获取某月某 Sheet 的数据（分页 + 搜索）。"""
    month = normalize_month(month)
    from bdms.modules.delivery_report.engine import DeliveryReportEngine
    engine = DeliveryReportEngine()
    data = engine.load(month)
    if sheet not in data:
        raise HTTPException(404, f"{month} 无 {sheet} 数据")
    df = data[sheet]
    
    # 搜索过滤（全字段模糊匹配）
    if search and search.strip():
        keyword = search.strip()
        mask = df.apply(
            lambda row: any(keyword.lower() in str(v).lower() for v in row.values),
            axis=1
        )
        df = df[mask]
    
    total = len(df)
    start = (page - 1) * page_size
    end = start + page_size
    page_df = df.iloc[start:end]
    return {
        "month": month,
        "sheet": sheet,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": (total + page_size - 1) // page_size,
        "columns": list(page_df.columns),
        "rows": page_df.fillna("").to_dict(orient="records"),
    }


# ========== 2. 确认收入 ==========

@router.get("/revenue/summary/{month}")
async def revenue_summary(month: str):
    month = normalize_month(month)
    from bdms.modules.revenue.summary_engine import SummaryEngine
    r = SummaryEngine().compute_summary(month)
    if "error" in r:
        raise HTTPException(400, r["error"])
    return r


@router.get("/revenue/export/{month}")
async def revenue_export(month: str):
    """导出确收汇总为 Excel。"""
    import io
    month = normalize_month(month)
    from bdms.modules.revenue.summary_engine import SummaryEngine
    r = SummaryEngine().compute_summary(month)
    if "error" in r:
        raise HTTPException(400, r["error"])
    
    try:
        import pandas as pd
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        wb = Workbook()
        
        # Sheet1: 汇总表
        ws = wb.active
        ws.title = "汇总表"
        
        rows_data = r.get("rows", [])
        if rows_data:
            # 表头
            headers = list(rows_data[0].keys())
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EDF5")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # 数据
            for row_idx, row in enumerate(rows_data, 2):
                for col_idx, key in enumerate(headers, 1):
                    ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
            
            # 列宽自适应
            for col in ws.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 30)
        
        # Sheet2: 月度趋势
        ws2 = wb.create_sheet("月度趋势")
        monthly = r.get("monthly", {})
        months = monthly.get("months", [])
        if months:
            ws2.cell(row=1, column=1, value="月份").font = Font(bold=True)
            for i, cat in enumerate(["新签", "递延", "合计"], 2):
                ws2.cell(row=1, column=i*2-1, value=f"{cat}-计划").font = Font(bold=True)
                ws2.cell(row=1, column=i*2, value=f"{cat}-实际").font = Font(bold=True)
            for row_idx, m in enumerate(months, 2):
                ws2.cell(row=row_idx, column=1, value=m)
                for i, cat in enumerate(["新签", "递延", "合计"], 2):
                    series = monthly.get(cat, {})
                    plan_val = series.get("plan", [])[row_idx-2] if row_idx-2 < len(series.get("plan", [])) else ""
                    actual_val = series.get("actual", [])[row_idx-2] if row_idx-2 < len(series.get("actual", [])) else ""
                    ws2.cell(row=row_idx, column=i*2-1, value=plan_val)
                    ws2.cell(row=row_idx, column=i*2, value=actual_val)
        
        # 保存到内存
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        from fastapi.responses import Response
        filename = f"确收汇总_{month}.xlsx"
        return Response(
            content=output.read(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=\"{filename}\""}
        )
    except Exception as e:
        raise HTTPException(500, f"导出失败: {e}")


@router.post("/revenue/generate")
async def revenue_generate(req: GenerateRequest):
    """生成/读取确认收入（幂等，受 mode 控制）。"""
    from bdms.modules.revenue.service import RevenueService
    svc = RevenueService()
    try:
        return svc.generate(normalize_month(req.month), getattr(req, "mode", "auto") or "auto")
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise HTTPException(500, str(e))


@router.post("/revenue/import")
async def revenue_import(req: GenerateRequest):
    from bdms.modules.revenue.service import RevenueService
    try:
        return RevenueService().import_source(normalize_month(req.month))
    except Exception as e:
        raise HTTPException(500, str(e))


@router.get("/revenue/compare/{month}")
async def revenue_compare(month: str, manual: Optional[str] = None):
    month = normalize_month(month)
    from pathlib import Path
    from bdms.core import paths
    from bdms.modules.revenue.summary_engine import SummaryEngine
    s = SummaryEngine()
    mine = s.compute_summary(month)
    if "error" in mine:
        raise HTTPException(400, mine["error"])
    manual_path = Path(manual) if manual else None
    if manual_path is None:
        # 在团队报告目录下查找该月的差异分析表
        mdir = paths.month_dir(month)
        if mdir.exists():
            hits = sorted(mdir.glob("*差异分析*.xlsx"))
            manual_path = hits[0] if hits else None
    if not manual_path or not Path(manual_path).exists():
        return {"mine": mine, "manual": None,
                "note": f"未找到 {month} 手工报表，仅返回本方计算"}
    manual = s.read_manual_summary(manual_path)
    mrows = {r.get("period"): r for r in manual.get("rows", []) if r.get("period")}
    compare = []
    for row in mine["rows"]:
        m = mrows.get(row["period"])
        in_scope = row["period"] <= month  # 同口径：手工报表月份及之前
        entry = {"period": row["period"], "categories": {}, "in_scope": in_scope}
        for cat, mkey in [("递延", "deferred"), ("新签", "new_sign"), ("合计", "total")]:
            op = row["categories"][cat]["plan"]
            oa = row["categories"][cat]["actual"]
            if m and in_scope:
                mp = (m[mkey]["plan"] or 0) * 10000
                ma = (m[mkey]["actual"] or 0) * 10000
                entry["categories"][cat] = {
                    "ours": {"plan": op, "actual": oa},
                    "manual": {"plan": mp, "actual": ma},
                    "match": abs(op - mp) < 1 and abs(oa - ma) < 1,
                }
            else:
                entry["categories"][cat] = {"ours": {"plan": op, "actual": oa}, "manual": None}
        compare.append(entry)
    ok = sum(1 for e in compare if e["in_scope"]
             for c in e["categories"].values() if c.get("match"))
    total = sum(1 for e in compare if e["in_scope"]
                for c in e["categories"].values() if c.get("manual") is not None)
    return {"mine": mine, "compare": compare, "matched": ok, "comparable": total,
            "all_match": ok == total and total > 0, "scope_month": month}


# ========== 3. 基础数据 ==========

@router.get("/master-data/types")
async def master_types():
    try:
        from bdms.modules.master_data.service import MasterDataService
        svc = MasterDataService()
        return {"types": svc.list_types()}
    except ImportError:
        raise HTTPException(503, "模块3（基础数据）尚未完成")


@router.get("/master-data/list/{data_type}")
async def master_list(data_type: str, include_disabled: bool = False):
    try:
        from bdms.modules.master_data.service import MasterDataService
        svc = MasterDataService()
        # 兼容不同命名：优先 list_items，回退 list_reference
        fn = getattr(svc, "list_items", None) or getattr(svc, "list_reference", None)
        if fn is None:
            raise HTTPException(500, "MasterDataService 缺少列表方法")
        try:
            items = fn(data_type, include_disabled=include_disabled)
        except TypeError:
            items = fn(data_type)
        return {"data_type": data_type, "items": items}
    except ImportError:
        raise HTTPException(503, "模块3（基础数据）尚未完成")


# ========== 4. 统计看板 ==========

@router.get("/dashboard/{month}")
async def dashboard(month: str, refresh: bool = False):
    try:
        from bdms.modules.dashboard.service import DashboardService
    except ImportError:
        raise HTTPException(503, "模块4（统计看板）尚未完成")
    svc = DashboardService()
    try:
        if refresh:
            svc.refresh_snapshot(month)
        return svc.get_dashboard(month)
    except Exception as e:
        raise HTTPException(500, str(e))


@router.post("/dashboard/drill-down")
async def dashboard_drill(req: DriftQuery):
    try:
        from bdms.modules.dashboard.service import DashboardService
    except ImportError:
        raise HTTPException(503, "模块4（统计看板）尚未完成")
    try:
        return DashboardService().drill_down(req.month, req.dimension, req.value,
                                             sheet=req.sheet)
    except Exception as e:
        raise HTTPException(500, str(e))


# ========== 5. 系统设定 ==========

@router.get("/settings")
async def settings_get():
    try:
        from bdms.modules.settings.service import SettingsService
        return SettingsService().get_all()
    except ImportError:
        # 降级：直接用 core
        return _db.get_all_settings()


@router.post("/settings")
async def settings_set(req: SettingUpdate):
    try:
        from bdms.modules.settings.service import SettingsService
        return SettingsService().set(req.key, req.value)
    except ImportError:
        _db.upsert_settings({req.key: req.value})
        return {"ok": True, "key": req.key, "value": req.value}


# ========== 健康检查 ==========

@router.get("/health")
async def health():
    return {"status": "ok", "component": "bdms-web"}
