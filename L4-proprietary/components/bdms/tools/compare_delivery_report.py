"""交付月报逐格比对工具（我方导出 vs 手工黄金基准）。

用法：
  python3 tools/compare_delivery_report.py 202606            # 全部 sheet
  python3 tools/compare_delivery_report.py 202606 签约 POC&提前实施
  python3 tools/compare_delivery_report.py 202606 --cols     # 附带逐列差异统计
"""
import sys, os, datetime as _dt
from collections import Counter
import openpyxl

MONTH = sys.argv[1] if len(sys.argv) > 1 and sys.argv[1].isdigit() else "202606"
ARGS = sys.argv[2:] if len(sys.argv) > 1 and sys.argv[1].isdigit() else sys.argv[1:]
SHOW_COLS = "--cols" in ARGS
SHEETS_ARG = [a for a in ARGS if not a.startswith("--")]

BASE = "/Users/bangcle/Bangcle Workspace/01. Management/2026/2026团队报告"
MINE = f"output/交付月报_{MONTH}.xlsx"
MAN = os.path.join(BASE, MONTH, f"2026交付月报-{MONTH}30.xlsx")
SHEETS = ["签约", "POC&提前实施", "异常项目", "确收交接", "验收交接"]


def norm(v):
    if v is None:
        return None
    if isinstance(v, _dt.time):
        return 0 if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0) else v
    if isinstance(v, _dt.datetime):
        # 手工把「日期」写成 datetime(2026,6,30,0,0)，我方写 date
        return v.date() if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0) else v
    if isinstance(v, str):
        s = v.strip()
        return None if s == "" else s
    return v


def eq(a, b):
    a, b = norm(a), norm(b)
    if a is None and b is None:
        return True
    if a is None or b is None:
        return False
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-6
    if isinstance(a, (int, float)) and isinstance(b, str):
        try: return abs(a - float(b)) <= 1e-6
        except Exception: return False
    if isinstance(b, (int, float)) and isinstance(a, str):
        try: return abs(float(a) - b) <= 1e-6
        except Exception: return False
    return str(a) == str(b)


def compare(sheet, wm=None, wa=None, detail=0):
    wm = wm or openpyxl.load_workbook(MINE, read_only=True, data_only=True)
    wa = wa or openpyxl.load_workbook(MAN, read_only=True, data_only=True)
    sm, sa = wm[sheet], wa[sheet]
    it_m, it_a = sm.iter_rows(values_only=True), sa.iter_rows(values_only=True)
    diffs, n, r = [], 0, 0
    bycol = Counter()
    while True:
        try: rowm = next(it_m)
        except StopIteration: rowm = None
        try: rowa = next(it_a)
        except StopIteration: rowa = None
        if rowm is None and rowa is None:
            break
        r += 1
        L = max(len(rowm or ()), len(rowa or ()))
        for c in range(L):
            vm = rowm[c] if rowm and c < len(rowm) else None
            va = rowa[c] if rowa and c < len(rowa) else None
            if not eq(vm, va):
                n += 1
                bycol[c + 1] += 1
                if len(diffs) < detail:
                    diffs.append((r, c + 1, vm, va))
    return n, diffs, r, sa.max_column, bycol


if __name__ == "__main__":
    wm = openpyxl.load_workbook(MINE, read_only=True, data_only=True)
    wa = openpyxl.load_workbook(MAN, read_only=True, data_only=True)
    sheets = SHEETS_ARG or SHEETS
    total = 0
    for sh in sheets:
        try:
            n, diffs, rows, cols, bycol = compare(sh, wm, wa, detail=10)
        except Exception as e:
            print(f"{sh}: ERROR {type(e).__name__}: {e}")
            continue
        total += n
        print(f"{sh}: {n} diffs  ({rows}r x {cols}c)")
        if SHOW_COLS and bycol:
            top = ", ".join(f"c{c}={k}" for c, k in bycol.most_common(20))
            print(f"    bycol: {top}")
        for (r, c, vm, va) in diffs[:6]:
            print(f"    r{r}c{c}: mine={vm!r} manual={va!r}")
    print(f"TOTAL: {total}")
