"""确收差异分析报表逐格比对工具。

用法：
  python3 tools/compare_report.py 202606              # 比对全部 10 个 sheet
  python3 tools/compare_report.py 202606 汇总 图例     # 只比对指定 sheet
"""
import sys, os, datetime as _dt
import openpyxl

_MONTH = sys.argv[1] if len(sys.argv) > 1 and sys.argv[1].isdigit() else "202606"
_SHEETS_ARG = [a for a in sys.argv[1:] if not a.isdigit()]
_D = "/Users/bangcle/Bangcle Workspace/01. Management/2026/2026团队报告"
MINE = f"output/确收差异分析_{_MONTH}.xlsx"
MAN = os.path.join(_D, _MONTH, "2026年计划确收&实际确收对比表202601-06-0724 - 差异分析.xlsx")

def norm(v):
    if v is None: return None
    # Excel 数值 0 在日期/时间格式列中会被 openpyxl 读成 time(0,0)
    if isinstance(v, _dt.time):
        return 0 if (v.hour, v.minute, v.second, v.microsecond) == (0,0,0,0) else v
    if isinstance(v, str):
        s = v.strip()
        return None if s == '' else s
    return v

def eq(a, b):
    a, b = norm(a), norm(b)
    if a is None and b is None: return True
    if a is None or b is None: return False
    if isinstance(a,(int,float)) and isinstance(b,(int,float)):
        return abs(a-b) <= 1e-6
    if isinstance(a,(int,float)) and isinstance(b,str):
        try: return abs(a-float(b)) <= 1e-6
        except: return False
    if isinstance(b,(int,float)) and isinstance(a,str):
        try: return abs(float(a)-b) <= 1e-6
        except: return False
    return str(a) == str(b)

def compare(sheet):
    wm = openpyxl.load_workbook(MINE, read_only=True, data_only=True)
    wa = openpyxl.load_workbook(MAN,  read_only=True, data_only=True)
    sm = wm[sheet]; sa = wa[sheet]
    maxr = max(sm.max_row, sa.max_row); maxc = max(sm.max_column, sa.max_column)
    diffs = []
    n = 0
    it_m = sm.iter_rows(values_only=True); it_a = sa.iter_rows(values_only=True)
    r = 0
    while True:
        try: rowm = next(it_m)
        except StopIteration: rowm = None
        try: rowa = next(it_a)
        except StopIteration: rowa = None
        if rowm is None and rowa is None: break
        r += 1
        L = max(len(rowm or ()), len(rowa or ()))
        for c in range(L):
            vm = rowm[c] if rowm and c < len(rowm) else None
            va = rowa[c] if rowa and c < len(rowa) else None
            if not eq(vm, va):
                n += 1
                if len(diffs) < 15:
                    diffs.append((r, c+1, vm, va))
    wm.close(); wa.close()
    return n, diffs, r, maxc

if __name__ == "__main__":
    sheets = _SHEETS_ARG or ["汇总","汇总分析","预算趋势分析","确收差异分析","预算执行表",
                             "计划确收底稿","重拆履约","图例","月度汇总记录","履约汇总记录"]
    for sh in sheets:
        try:
            n, diffs, rows, cols = compare(sh)
            print(f"{sh}: {n} diffs  ({rows}r x {cols}c)")
            for (r, c, vm, va) in diffs[:8]:
                print(f"    r{r}c{c}: mine={vm!r} manual={va!r}")
        except Exception as e:
            print(f"{sh}: ERROR {e}")
