"""交付月报模块 v2.1 — 单元测试。

对齐 DESIGN-DETAIL-DELIVERY-REPORT-v2.1.md：
  1. Validator 12 条规则（非空/类型/枚举/行数/重复）
  2. 存疑数据收集 + 人工调整
  3. DASHBOARD 连接器（统计 Sheet 聚合）
  4. 图例预落盘 + 读取
  5. 15 Sheet 导出（结构）
"""

import sys
import pytest
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT / "src") not in sys.path:
    sys.path.insert(0, str(_ROOT / "src"))

import pandas as pd

from bdms.core import db as _db
from bdms.modules.delivery_report.validator import DeliveryReportValidator
from bdms.modules.dashboard.delivery_report_connector import DeliveryReportConnector


# ─── Fixtures ───

@pytest.fixture
def db_path(tmp_path):
    path = tmp_path / "test.db"
    _db.init_db(path)
    return path


@pytest.fixture
def validator(db_path):
    return DeliveryReportValidator(db_path=db_path)


# ─── 1. Validator 规则 ───

class TestValidatorRules:

    def test_not_null(self, validator):
        df = pd.DataFrame({
            "合同编号": ["A", None, ""],
            "项目名称": ["x", "y", "z"],
        })
        errors = validator.validate_not_null(df, ["合同编号"])
        assert len(errors) == 2
        assert errors[0]["row"] == 1 and errors[1]["row"] == 2

    def test_not_null_missing_column(self, validator):
        df = pd.DataFrame({"a": [1]})
        errors = validator.validate_not_null(df, ["不存在的列"])
        assert len(errors) == 1
        assert errors[0]["issue"] == "missing_column"

    def test_date_type(self, validator):
        df = pd.DataFrame({"日期": ["2026-06-01", "2026/7/1", "bogus", ""]})
        errors = validator.validate_data_type(df, {"日期": "date"})
        # bogus 是非法日期；空值跳过
        assert len(errors) == 1
        assert errors[0]["row"] == 2

    def test_numeric_type(self, validator):
        df = pd.DataFrame({"金额": ["100", "200.5", "abc", "-3"]})
        errors = validator.validate_data_type(df, {"金额": "numeric"})
        assert len(errors) == 1 and errors[0]["row"] == 2

    def test_enum(self, validator):
        df = pd.DataFrame({"状态": ["正常", "异常", "未知"]})
        errors = validator.validate_enum(df, {"状态": ["正常", "异常"]})
        assert len(errors) == 1 and errors[0]["value"] == "未知"

    def test_row_count_within_threshold(self, validator):
        r = validator.validate_row_count("测试", 100, 110, threshold=0.2)
        assert not r["warning"]

    def test_row_count_exceeds_threshold(self, validator):
        r = validator.validate_row_count("测试", 100, 200, threshold=0.2)
        assert r["warning"] and r["deviation"] == 0.5

    def test_row_count_no_previous(self, validator):
        r = validator.validate_row_count("测试", 100, 0)
        assert not r["warning"]

    def test_duplicates(self, validator):
        df = pd.DataFrame({"合同编号": ["A", "B", "A"]})
        errors = validator.validate_duplicates(df, ["合同编号"])
        assert len(errors) == 2  # 两行都标记


# ─── 2. 存疑数据 + 人工调整 ───

class TestSuspiciousFlow:

    def _seed_month(self, db_path, month, rows):
        """直接写 dr_sheet_row 测试数据。"""
        import json
        conn = _db.get_connection(db_path)
        try:
            columns = ["合同编号", "项目名称", "签约日期"]
            conn.execute(
                "DELETE FROM dr_sheet_row WHERE month = ?", (month,))
            for i, r in enumerate(rows):
                conn.execute(
                    "INSERT INTO dr_sheet_row (month, sheet, row_index, data) "
                    "VALUES (?, '签约', ?, ?)",
                    (month, i, json.dumps(r, ensure_ascii=False)),
                )
            conn.execute(
                "INSERT OR REPLACE INTO dr_sheet_meta (month, sheet, columns, row_count) "
                "VALUES (?, '签约', ?, ?)",
                (month, json.dumps(columns), len(rows)),
            )
            conn.commit()
        finally:
            conn.close()

    def test_collect_suspicious(self, db_path, validator):
        self._seed_month(db_path, "202606", [
            {"合同编号": "A", "项目名称": "p1", "签约日期": "2026-06-01"},
            {"合同编号": "", "项目名称": "p2", "签约日期": "bad-date"},
        ])
        result = validator.collect_suspicious("202606")
        assert result["summary"]["total_errors"] >= 2  # 空编号 + 非法日期
        assert any(e["issue"] == "null" for e in result["errors"])
        assert any(e["issue"] == "type_date" for e in result["errors"])

    def test_apply_manual_fixes(self, db_path, validator):
        self._seed_month(db_path, "202606", [
            {"合同编号": "", "项目名称": "p1"},
        ])
        result = validator.apply_manual_fixes("202606", [
            {"sheet": "签约", "row_index": 0, "col": "合同编号", "value": "FIXED-001"},
        ])
        assert result["applied"] == 1

        # 验证修复生效
        import json
        conn = _db.get_connection(db_path)
        row = conn.execute(
            "SELECT data FROM dr_sheet_row WHERE month='202606' AND sheet='签约' AND row_index=0"
        ).fetchone()
        conn.close()
        data = json.loads(row[0])
        assert data["合同编号"] == "FIXED-001"


# ─── 3. DASHBOARD 连接器 ───

class TestConnector:

    def test_load_empty_month(self, db_path):
        c = DeliveryReportConnector(db_path)
        assert c.build_stats_sheets("209901") is not None

    def test_abnormal_ledger_empty(self, db_path):
        c = DeliveryReportConnector(db_path)
        df = c.build_abnormal_ledger("209901")
        # 无数据时返回空 DataFrame（不报错）
        assert df is not None

    def test_legend_config_empty(self, db_path):
        c = DeliveryReportConnector(db_path)
        df = c.get_legend_config()
        assert df is not None  # 空库不报错


# ─── 4. 图例预落盘（真实黄金基准）───

class TestLegendSeed:

    def test_seed_legend_idempotent(self):
        """预落盘幂等（重复执行不报错不重复）。"""
        from bdms.modules.delivery_report.seed_legend import seed_legend
        import tempfile
        from pathlib import Path
        tmp = Path(tempfile.mkdtemp()) / "t.db"
        _db.init_db(tmp)
        r1 = seed_legend(db_path=tmp)
        assert r1["seeded"] > 0
        r2 = seed_legend(db_path=tmp)
        assert r2["seeded"] == r1["seeded"]  # 幂等

        conn = _db.get_connection(tmp)
        n = conn.execute(
            "SELECT COUNT(*) FROM md_reference WHERE data_type='legend_config'"
        ).fetchone()[0]
        conn.close()
        assert n == r1["seeded"]


# ─── 5. 导出结构（真实数据，慢速标记）───

@pytest.mark.slow
class TestExportStructure:
    """真实数据导出测试（需要 ONES 数据源 + 已生成 DB）。

    标记 slow，全量跑时用 -m 'not slow' 跳过。
    """

    def test_export_15_sheets(self):
        from bdms.modules.delivery_report.exporter import DeliveryReportExporter
        e = DeliveryReportExporter()
        path = e.export("202606", out_path="/tmp/bdms_test_export.xlsx")
        assert path.exists()

        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        assert len(wb.sheetnames) == 15
        assert wb.sheetnames[0] == "签约"
        assert wb.sheetnames[-1] == "图例"
        wb.close()
