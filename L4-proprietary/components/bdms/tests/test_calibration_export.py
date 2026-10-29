"""校准导出增强 — 单元测试。

对齐 DESIGN-DETAIL-MANUAL-CALIBRATION-v1.0.md §4.7：
  - 核心 Sheet 数据来自 effective（校准覆盖层唯一口径）
  - 已校准单元格带批注（原值→新值 / 操作人 / 时间）
  - 末尾追加「校准台账」Sheet（= ledger_rows）
统计聚合 Sheet（build_stats_sheets）与图例非本任务范围，monkeypatch 隔离。
"""
import sys
from pathlib import Path

import pandas as pd
import pytest

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT / "src") not in sys.path:
    sys.path.insert(0, str(_ROOT / "src"))

from bdms.core import db as _db
from bdms.modules.delivery_report.calibration import CalibrationService
from bdms.modules.delivery_report.engine import DeliveryReportEngine
from bdms.modules.delivery_report.exporter import DeliveryReportExporter


@pytest.fixture
def db_path(tmp_path):
    path = tmp_path / "test.db"
    _db.init_db(path)
    return path


@pytest.fixture
def seeded(db_path):
    raw = {
        "签约": pd.DataFrame({
            "合同编号": ["C001", "C002"],
            "项目名称": ["P-A", "P-B"],
            "所属产线": ["L1", "L2"],
            "签约日期": ["2026/6/3", "bad-date"],
            "金额": ["100", "200"],
        }),
    }
    DeliveryReportEngine(db_path).persist("202610", raw)
    return db_path


@pytest.fixture
def exporter_with_stats_isolated(db_path, monkeypatch):
    """导出器 + 隔离统计聚合与 SUBTOTAL 公式（避免依赖全量业务列布局）。"""
    from bdms.modules.dashboard.delivery_report_connector import DeliveryReportConnector
    import bdms.modules.delivery_report.exporter as exporter_mod
    monkeypatch.setattr(
        DeliveryReportConnector, "build_stats_sheets",
        lambda self, month: {})
    monkeypatch.setattr(
        DeliveryReportConnector, "get_legend_config",
        lambda self: None)
    # SUBTOTAL 公式列依赖真实 66 列布局，与本任务（校准覆盖/批注/台账）无关
    monkeypatch.setattr(exporter_mod, "SUBTOTAL_FORMULA_COLS", {})
    return DeliveryReportExporter(db_path)


def _export(exp, month="202610"):
    out = exp.export(month, out_path=exp.db_path.parent / "report.xlsx")
    import openpyxl
    return out, openpyxl.load_workbook(out)


def test_export_uses_effective_values(exporter_with_stats_isolated, seeded):
    CalibrationService(seeded).upsert(
        "202610", "签约", 1, "签约日期", "2026-06-05", "日期修正", "gateway-owner")
    _, wb = _export(exporter_with_stats_isolated)
    ws = wb["签约"]
    # 标题 r1、表头 r2、数据从 r3；0-based 行1 → Excel r4
    assert ws.cell(row=4, column=4).value == "2026-06-05"
    # 未校准行保持原值
    assert ws.cell(row=3, column=4).value == "2026/6/3"


def test_export_adds_comment_on_calibrated_cell(exporter_with_stats_isolated, seeded):
    CalibrationService(seeded).upsert(
        "202610", "签约", 1, "签约日期", "2026-06-05", "日期修正", "gateway-owner")
    _, wb = _export(exporter_with_stats_isolated)
    ws = wb["签约"]
    cm = ws.cell(row=4, column=4).comment
    assert cm is not None
    assert "2026-06-05" in cm.text
    assert "gateway-owner" in cm.text
    # 未校准单元格无批注
    assert ws.cell(row=3, column=4).comment is None


def test_export_appends_ledger_sheet(exporter_with_stats_isolated, seeded):
    CalibrationService(seeded).upsert(
        "202610", "签约", 1, "签约日期", "2026-06-05", "原因X", "gateway-owner")
    _, wb = _export(exporter_with_stats_isolated)
    assert "校准台账" in wb.sheetnames
    led = wb["校准台账"]
    assert led.max_row >= 2
    assert led.cell(row=1, column=1).value == "月份"
    assert led.cell(row=2, column=2).value == "签约"
    assert led.cell(row=2, column=6).value == "2026-06-05"
    assert led.cell(row=2, column=9).value == "原因X"


def test_export_without_calibration_omits_ledger(exporter_with_stats_isolated, seeded):
    _, wb = _export(exporter_with_stats_isolated)
    assert "校准台账" not in wb.sheetnames  # 无 active 校准 → 省略台账 Sheet


def test_export_original_cells_unaffected_by_comment(exporter_with_stats_isolated, seeded):
    """导出后原始 DB 数据仍不被改动（effective 只在导出层）。"""
    CalibrationService(seeded).upsert(
        "202610", "签约", 1, "签约日期", "2026-06-05", "日期修正", "gateway-owner")
    _, wb = _export(exporter_with_stats_isolated)
    conn = _db.get_connection(seeded)
    try:
        import json
        row = json.loads(conn.execute(
            "SELECT data FROM dr_sheet_row WHERE month='202610' "
            "AND sheet='签约' AND row_index=1").fetchone()["data"])
    finally:
        conn.close()
    assert row["签约日期"] == "bad-date"  # 原始层不变
